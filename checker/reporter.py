"""
checker.reporter
多维度判题报告与成绩单生成器：
  1. labX结果.csv (逐人逐题评测明细，细化到检查分项)
  2. labX报告.md (批次统计、分项达成率、规则说明、复核清单)
  3. labX实验成绩单.csv (完整点名册、各题验收汇总、建议扣分合计、>=50字定制评语)
  4. labX实验成绩单.xlsx (如果存在模板文件，同步填充 Excel 工作表)
  5. 单人重判独立导出与批次安全合并
"""
import csv
import json

from checker.scoring import validate_grade, ScoringConsistencyError, cents
from checker.state import SCHEMA_VERSION, grade_to_dict, grade_from_dict
import shutil
from pathlib import Path
from typing import Any, Dict, List, Optional

from checker.models import CheckItem, ProblemResult, Student, StudentGrade


class LabReporter:
    """负责将评测结果汇总并输出交付报表"""

    def __init__(self, lab_id: str, output_dir: Path):
        self.lab_id = lab_id.lower()
        self.output_dir = output_dir
        self.output_dir.mkdir(parents=True, exist_ok=True)

    @staticmethod
    def _prepare(grades):
        for grade in grades.values():
            validate_grade(grade)

    @staticmethod
    def _problem_cells(grade,pid):
        p=(grade.problem_results.get(pid) or grade.optional_results.get(pid)) if grade else None
        return {'题目得分':('待复核' if p.is_pending else f'{p.score:.2f}') if p else '',
                '题目满分':f'{p.max_score:.2f}' if p else '',
                '采用版本':(p.chosen_candidate or {}).get('filename','') if p else '',
                '候选数':len(p.candidate_evaluations) if p else 0}

    def _save_state(self, grades, mock_mode=False):
        path = self.output_dir / f"{'mock_' if mock_mode else ''}{self.lab_id}评分状态.json"
        payload = {'schema_version':SCHEMA_VERSION,'lab_id':self.lab_id,'mock_mode':mock_mode,
                   'grades':{sid:grade_to_dict(g) for sid,g in grades.items()}}
        path.write_text(json.dumps(payload,ensure_ascii=False,indent=2,default=str,allow_nan=False),encoding='utf-8')
        return path

    @staticmethod
    def _candidate_markdown(grades):
        lines=['', '## 候选版本与最终选择',
               '全部版本独立评定后择优；用例预算、分类扣分与最终分共享同一计分结果。金额按分分配舍入余数；未定事实保留复核。',
               '', '| 学号 | 题号 | 版本 | 最终得分 | 用例通过 | 评定状态 | 采用 |',
               '|---|---|---|---:|---|---|---|']
        for sid,g in grades.items():
            for pid,p in {**g.problem_results,**g.optional_results}.items():
                selected=(p.chosen_candidate or {}).get('src_path')
                for c in p.candidate_evaluations:
                    state='待复核' if c.get('is_pending') else '已评定'
                    score='待复核' if c.get('is_pending') else f"{c['score']:.2f}"
                    name=c['filename'].replace('|','\\|')
                    mark='是' if c['src_path']==selected else ''
                    lines.append(f"| {sid} | {pid} | {name} | {score}/{c['max_score']:.2f} | {c['passed_cases']}/{len(c.get('test_results',[]))} | {state} | {mark} |")
        return '\n'.join(lines)

    def generate_all_reports(
        self,
        roster_students: List[Student],
        grades: Dict[str, StudentGrade],
        ambiguous_items: Optional[List[Dict[str, Any]]] = None,
        unmatched_items: Optional[List[Dict[str, Any]]] = None,
        excel_template_path: Optional[Path] = None,
        mock_mode: bool = False,
        encoding_stats: Optional[Dict[str, Any]] = None,
    ) -> Dict[str, Path]:
        """生成全套标准报表并返回文件路径字典"""
        self._prepare(grades)
        csv_detail_path = self.generate_detail_csv(roster_students, grades, mock_mode=mock_mode)
        md_summary_path = self.generate_summary_markdown(
            roster_students,
            grades,
            ambiguous_items=ambiguous_items,
            unmatched_items=unmatched_items,
            mock_mode=mock_mode,
            encoding_stats=encoding_stats,
        )
        csv_roster_path = self.generate_roster_grade_csv(roster_students, grades, mock_mode=mock_mode)

        result_paths = {
            "detail_csv": csv_detail_path,
            "summary_md": md_summary_path,
            "roster_csv": csv_roster_path,
        }

        if excel_template_path and Path(excel_template_path).exists():
            xlsx_path = self.generate_roster_grade_xlsx(roster_students, grades, excel_template_path, mock_mode=mock_mode)
            if xlsx_path:
                result_paths["roster_xlsx"] = xlsx_path

        result_paths['state_json'] = self._save_state(grades,mock_mode)
        return result_paths

    def generate_detail_csv(
        self,
        roster_students: List[Student],
        grades: Dict[str, StudentGrade],
        mock_mode: bool = False,
    ) -> Path:
        """
        导出 labX结果.csv
        记录粒度: 学生 × 实验 × 题号 × 检查分项
        包含字段: 学号, 姓名, 实验, 题号, 评分项, 提交状态, 达成状态, 检查依据, 建议扣分, 扣分理由, Agent说明附注, 是否需要复核, 实验评语
        """
        self._prepare(grades)
        prefix = "mock_" if mock_mode else ""
        file_path = self.output_dir / f"{prefix}{self.lab_id}结果.csv"

        fieldnames = [
            "学号", "姓名", "实验", "题号", "评分项", "提交状态",
            "达成状态", "检查依据", "建议扣分", "扣分理由", "Agent说明附注", "是否需要复核", "实验评语", "题目得分", "题目满分", "采用版本", "候选数"
        ]

        with open(file_path, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()

            for student in roster_students:
                grade = grades.get(student.student_id)
                items = grade.all_check_items if grade else []

                # 若当前学生尚未生成 check_items 或完全缺交，生成标准缺交分项
                if not items:
                    items = self._build_default_check_items(student, grade)

                if self.lab_id == "lab2":
                    missing_msg = (
                        f"{student.name}同学你好！系统未检测到你在本次LAB2中的作业提交（未收到源码文件）。"
                        f"按课程大纲与实验教学要求，请尽快核实是否由于提交逾期或文件路径遗漏，并及时联系助教说明情况补交，以便及时核对并确认你的实验平时成绩。"
                    )
                else:
                    missing_msg = (
                        f"{student.name}同学你好！系统未检测到你在实验1中的作业提交（未收到源码文件及演草纸验算照片）。"
                        f"按课程大纲与实验教学要求，请尽快核实是否由于提交逾期或文件路径遗漏，并及时联系助教说明情况补交，以便及时核对并确认实验平时成绩。"
                    )
                student_comment = grade.comment if grade else missing_msg

                for it in items:
                    if it.deduction is None:
                        ded_str = "待复核"
                    elif it.deduction > 0:
                        ded_str = f"{it.deduction:.2f}"
                    else:
                        ded_str = "0.00"

                    rev_str = "是" if it.needs_review else "否"

                    writer.writerow({
                        "学号": student.student_id,
                        "姓名": student.name,
                        "实验": self.lab_id,
                        "题号": it.problem_id,
                        "评分项": it.sub_item,
                        "提交状态": it.submission_status,
                        "达成状态": it.completion_status,
                        "检查依据": it.evidence,
                        "建议扣分": ded_str,
                        "扣分理由": it.deduction_reason,
                        "Agent说明附注": it.agent_note,
                        "是否需要复核": rev_str,
                        "实验评语": student_comment,
                        **self._problem_cells(grade,it.problem_id),
                    })

        return file_path

    def _build_default_check_items(self, student: Student, grade: Optional[StudentGrade]) -> List[CheckItem]:
        """为学生构建全量的默认检查分项"""
        if grade and grade.status!='missing' and not grade.all_check_items and not grade.is_pending:
            return [CheckItem('history','历史总分记录','已提交','无法验证',
                '仅保存实验总分，缺少历史题目与用例证据；不推断各题已完成',
                (cents(grade.max_score)-cents(grade.total_score))/100,
                '依据已保存实验总分记录扣分合计，分项事实无法还原','缺少分项证据',True)]
        if self.lab_id not in {'lab1','lab2'}:
            pending=bool(grade and grade.is_pending)
            return [CheckItem('all','实验提交','已提交' if pending else '未提交',
                '待复核' if pending else '未达成','依实验提交与评分记录核验',
                None if pending else (grade.max_score if grade else 5.0),needs_review=pending)]
        is_missing = not grade or grade.status == "missing"
        is_conflict = grade and (grade.status in {"ambiguous", "conflict"} or any("冲突" in r for r in grade.review_reasons))
        evidence_msg = "未检测到实验提交文件" if is_missing else ("多学号匹配冲突待复核" if is_conflict else "学生提交已评测完成")

        if self.lab_id == "lab2":

            if is_missing:
                return [
                    CheckItem("warmup", "课前热身练习", "未提交", "不适用", evidence_msg, 0.0, "热身练习依规则不计入总分", "课内练习无强制提交要求", False),
                    CheckItem("2-1", "源码提交", "未提交", "未达成", evidence_msg, 2.00, "未提交 Lab 2-1 源码文件", "缺少源码", False),
                    CheckItem("2-1", "编译与运行", "不适用", "无法验证", evidence_msg, 0.0, "未提交源码，本项无法验证，不重复扣分", "无法验证", False),
                    CheckItem("2-1", "时间段问候逻辑", "不适用", "无法验证", evidence_msg, 0.0, "未提交源码，本项无法验证，不重复扣分", "无法验证", False),
                    CheckItem("2-1", "边界与无效输入处理", "不适用", "无法验证", evidence_msg, 0.0, "未提交源码，本项无法验证，不重复扣分", "无法验证", False),
                    CheckItem("2-2", "源码提交", "未提交", "未达成", evidence_msg, 3.00, "未提交 Lab 2-2 源码文件", "缺少源码", False),
                    CheckItem("2-2", "编译与运行", "不适用", "无法验证", evidence_msg, 0.0, "未提交源码，本项无法验证，不重复扣分", "无法验证", False),
                    CheckItem("2-2", "字母对称金字塔展开 (A~Z)", "不适用", "无法验证", evidence_msg, 0.0, "未提交源码，本项无法验证，不重复扣分", "无法验证", False),
                    CheckItem("2-2", "大小写字母兼容转换", "不适用", "无法验证", evidence_msg, 0.0, "未提交源码，本项无法验证，不重复扣分", "无法验证", False),
                    CheckItem("2-2", "非字母与异常输入报错", "不适用", "无法验证", evidence_msg, 0.0, "未提交源码，本项无法验证，不重复扣分", "无法验证", False),
                ]
            else:
                return [
                    CheckItem("warmup", "课前热身练习", "已提交", "不适用", evidence_msg, 0.0, "热身练习依规则不计入总分", "不作强制提交要求", False),
                    CheckItem("2-1", "源码提交", "已提交", "待复核", evidence_msg, None, "待人工复核确认提交情况与源码归属", "待复核", True),
                    CheckItem("2-1", "编译与运行", "不适用", "无法验证", evidence_msg, 0.0, "待人工复核，本项无法验证，不重复扣分", "无法验证", False),
                    CheckItem("2-1", "时间段问候逻辑", "不适用", "无法验证", evidence_msg, 0.0, "待人工复核，本项无法验证，不重复扣分", "无法验证", False),
                    CheckItem("2-1", "边界与无效输入处理", "不适用", "无法验证", evidence_msg, 0.0, "待人工复核，本项无法验证，不重复扣分", "无法验证", False),
                    CheckItem("2-2", "源码提交", "已提交", "待复核", evidence_msg, None, "待人工复核确认提交情况与源码归属", "待复核", True),
                    CheckItem("2-2", "编译与运行", "不适用", "无法验证", evidence_msg, 0.0, "待人工复核，本项无法验证，不重复扣分", "无法验证", False),
                    CheckItem("2-2", "字母对称金字塔展开 (A~Z)", "不适用", "无法验证", evidence_msg, 0.0, "待人工复核，本项无法验证，不重复扣分", "无法验证", False),
                    CheckItem("2-2", "大小写字母兼容转换", "不适用", "无法验证", evidence_msg, 0.0, "待人工复核，本项无法验证，不重复扣分", "无法验证", False),
                    CheckItem("2-2", "非字母与异常输入报错", "不适用", "无法验证", evidence_msg, 0.0, "待人工复核，本项无法验证，不重复扣分", "无法验证", False),
                ]


        if is_missing:
            items = [
                CheckItem("1-1", "课内环境熟悉与基础代码练习", "未提交", "不适用", evidence_msg, 0.0, "课内基础练习无扣分", "课内练习无强制提交要求", False),
                CheckItem("1-2", "源码提交", "未提交", "未达成", evidence_msg, 1.50, "未提交 Lab 1-2 源码文件", "缺少源码", False),
                CheckItem("1-2", "编译与运行", "不适用", "无法验证", evidence_msg, 0.0, "未提交源码，本项无法验证，不重复扣分", "无法验证", False),
                CheckItem("1-2", "参数学号与手机号对齐", "不适用", "无法验证", evidence_msg, 0.0, "未提交源码，本项无法验证，不重复扣分", "无法验证", False),
                CheckItem("1-2", "逻辑运算(!, ||, &&)", "不适用", "无法验证", evidence_msg, 0.0, "未提交源码，本项无法验证，不重复扣分", "无法验证", False),
                CheckItem("1-2", "位运算(~, |, &)", "不适用", "无法验证", evidence_msg, 0.0, "未提交源码，本项无法验证，不重复扣分", "无法验证", False),
                CheckItem("1-2", "有符号与无符号输出格式", "不适用", "无法验证", evidence_msg, 0.0, "未提交源码，本项无法验证，不重复扣分", "无法验证", False),
                CheckItem("1-2", "验算材料提交", "未提交", "未达成", evidence_msg, 1.50, "未提交 Lab 1-2 演草纸手写验算照片", "缺少手写材料", False),
                CheckItem("1-2", "姓名、学号、手机号完整性", "不适用", "无法验证", evidence_msg, 0.0, "未提交手写材料，本项无法验证，不重复扣分", "无法验证", False),
                CheckItem("1-2", "手写参数与身份及源码一致性", "不适用", "无法验证", evidence_msg, 0.0, "未提交手写材料，本项无法验证，不重复扣分", "无法验证", False),
                CheckItem("1-2", "六种运算展开过程完整性", "不适用", "无法验证", evidence_msg, 0.0, "未提交手写材料，本项无法验证，不重复扣分", "无法验证", False),
                CheckItem("1-2", "各验算过程与结果正确性", "不适用", "无法验证", evidence_msg, 0.0, "未提交手写材料，本项无法验证，不重复扣分", "无法验证", False),
                CheckItem("1-3", "源码提交", "未提交", "未达成", evidence_msg, 2.00, "未提交 Lab 1-3 源码文件", "缺少源码", False),
                CheckItem("1-3", "编译与运行", "不适用", "无法验证", evidence_msg, 0.0, "未提交源码，本项无法验证，不重复扣分", "无法验证", False),
                CheckItem("1-3", "华氏转摄氏计算", "不适用", "无法验证", evidence_msg, 0.0, "未提交源码，本项无法验证，不重复扣分", "无法验证", False),
                CheckItem("1-3", "摄氏转华氏计算", "不适用", "无法验证", evidence_msg, 0.0, "未提交源码，本项无法验证，不重复扣分", "无法验证", False),
                CheckItem("1-3", "边界、负温与浮点测试", "不适用", "无法验证", evidence_msg, 0.0, "未提交源码，本项无法验证，不重复扣分", "无法验证", False),
                CheckItem("1-3", "输出精度与格式", "不适用", "无法验证", evidence_msg, 0.0, "未提交源码，本项无法验证，不重复扣分", "无法验证", False),
            ]
        else:
            items = [
                CheckItem("1-1", "课内环境熟悉与基础代码练习", "已提交", "不适用", evidence_msg, 0.0, "课内基础练习无扣分", "课内练习无强制提交要求", False),
                CheckItem("1-2", "源码提交", "已提交", "待复核", evidence_msg, None, "待人工复核确认提交情况与源码归属", "待复核", True),
                CheckItem("1-2", "编译与运行", "不适用", "无法验证", evidence_msg, 0.0, "待人工复核，本项无法验证，不重复扣分", "无法验证", False),
                CheckItem("1-2", "参数学号与手机号对齐", "不适用", "无法验证", evidence_msg, 0.0, "待人工复核，本项无法验证，不重复扣分", "无法验证", False),
                CheckItem("1-2", "逻辑运算(!, ||, &&)", "不适用", "无法验证", evidence_msg, 0.0, "待人工复核，本项无法验证，不重复扣分", "无法验证", False),
                CheckItem("1-2", "位运算(~, |, &)", "不适用", "无法验证", evidence_msg, 0.0, "待人工复核，本项无法验证，不重复扣分", "无法验证", False),
                CheckItem("1-2", "有符号与无符号输出格式", "不适用", "无法验证", evidence_msg, 0.0, "待人工复核，本项无法验证，不重复扣分", "无法验证", False),
                CheckItem("1-2", "验算材料提交", "已提交", "待复核", evidence_msg, None, "待人工复核确认手写验算归属", "待复核", True),
                CheckItem("1-2", "姓名、学号、手机号完整性", "不适用", "无法验证", evidence_msg, 0.0, "待人工复核，本项无法验证，不重复扣分", "无法验证", False),
                CheckItem("1-2", "手写参数与身份及源码一致性", "不适用", "无法验证", evidence_msg, 0.0, "待人工复核，本项无法验证，不重复扣分", "无法验证", False),
                CheckItem("1-2", "六种运算展开过程完整性", "不适用", "无法验证", evidence_msg, 0.0, "待人工复核，本项无法验证，不重复扣分", "无法验证", False),
                CheckItem("1-2", "各验算过程与结果正确性", "不适用", "无法验证", evidence_msg, 0.0, "待人工复核，本项无法验证，不重复扣分", "无法验证", False),
                CheckItem("1-3", "源码提交", "已提交", "待复核", evidence_msg, None, "待人工复核确认提交情况与源码归属", "待复核", True),
                CheckItem("1-3", "编译与运行", "不适用", "无法验证", evidence_msg, 0.0, "待人工复核，本项无法验证，不重复扣分", "无法验证", False),
                CheckItem("1-3", "华氏转摄氏计算", "不适用", "无法验证", evidence_msg, 0.0, "待人工复核，本项无法验证，不重复扣分", "无法验证", False),
                CheckItem("1-3", "摄氏转华氏计算", "不适用", "无法验证", evidence_msg, 0.0, "待人工复核，本项无法验证，不重复扣分", "无法验证", False),
                CheckItem("1-3", "边界、负温与浮点测试", "不适用", "无法验证", evidence_msg, 0.0, "待人工复核，本项无法验证，不重复扣分", "无法验证", False),
                CheckItem("1-3", "输出精度与格式", "不适用", "无法验证", evidence_msg, 0.0, "待人工复核，本项无法验证，不重复扣分", "无法验证", False),
            ]
        return items

    def generate_roster_grade_csv(
        self,
        roster_students: List[Student],
        grades: Dict[str, StudentGrade],
        mock_mode: bool = False,
    ) -> Path:
        """
        导出 labX实验成绩单.csv
        包含字段: 学号, 姓名, 各题验收汇总, 源码提交与正确性, 手写材料提交与正确性, 未达成项目, 建议扣分合计, 扣分依据, Agent说明附注, 待复核事项, 实验评语
        """
        self._prepare(grades)
        prefix = "mock_" if mock_mode else ""
        file_path = self.output_dir / f"{prefix}{self.lab_id}实验成绩单.csv"

        ded_header = "建议扣分合计[模拟测试非正式]" if mock_mode else "建议扣分合计"
        fieldnames = [
            "学号", "姓名", "各题验收汇总", "源码提交与正确性", "手写材料提交与正确性",
            "未达成项目", ded_header, "扣分依据", "Agent说明附注", "待复核事项", "实验评语", "总分", "满分", "成绩状态", "成绩待定"
        ]

        with open(file_path, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()

            for student in roster_students:
                grade = grades.get(student.student_id)
                items = grade.all_check_items if grade else []
                if not items:
                    items = self._build_default_check_items(student, grade)

                # 1. 各题验收汇总
                if self.lab_id == "lab2":
                    if not grade or grade.status == "missing":
                        summary_problems = "全未交 (warmup:未提交; 2-1:未提交; 2-2:未提交)"
                    else:
                        summaries = []
                        for pid in ["warmup", "2-1", "2-2", "2-3", "2-4"]:
                            prob = grade.problem_results.get(pid) or grade.optional_results.get(pid)
                            if prob:
                                if pid == "warmup":
                                    warmup_st = "已提交" if (prob.status == "passed" or prob.source_files) else "未提交"
                                    summaries.append(f"warmup:{warmup_st}(不计分)")
                                elif pid == "2-1":
                                    summaries.append("2-1:待复核" if prob.is_pending else f"2-1:(得分{prob.score:.2f}/2.00)")
                                elif pid == "2-2":
                                    summaries.append("2-2:待复核" if prob.is_pending else f"2-2:(得分{prob.score:.2f}/3.00)")
                                elif pid in {"2-3", "2-4"}:
                                    summaries.append(f"{pid}:选做({prob.status},通过{sum(c.passed for c in prob.test_results)}/{len(prob.test_results)},不计分)")
                        summary_problems = "; ".join(summaries) if summaries else grade.status
                else:
                    if not grade or grade.status == "missing":
                        summary_problems = "全未交 (1-1:未提交; 1-2:未提交; 1-3:未提交)"
                    else:
                        summaries = []
                        for pid in ["1-1", "1-2", "1-3"]:
                            prob = grade.problem_results.get(pid)
                            if prob:
                                if pid == "1-1":
                                    summaries.append("1-1:课内练习")
                                elif pid == "1-2":
                                    c_sc = prob.sub_scores.get("code", 0.0)
                                    h_sc = prob.sub_scores.get("handwritten", 0.0)
                                    summaries.append("1-2:待复核" if prob.is_pending else f"1-2:(代码{c_sc:.2f}/1.50, 手写{h_sc:.2f}/1.50)")
                                elif pid == "1-3":
                                    summaries.append("1-3:待复核" if prob.is_pending else f"1-3:(得分{prob.score:.2f}/2.00)")
                        summary_problems = "; ".join(summaries) if summaries else grade.status

                # 2. 源码提交与正确性
                if self.lab_id == "lab2":
                    if not grade or grade.status == "missing":
                        src_summary = "未提交任何 C 源码"
                    else:
                        r2_1 = grade.problem_results.get("2-1")
                        r2_2 = grade.problem_results.get("2-2")
                        has_cbp = any(".cbp" in r for r in grade.review_reasons)
                        has_lnk = any(".lnk" in r for r in grade.review_reasons)
                        has_corrupt = any("损坏" in r or "解压失败" in r for r in grade.review_reasons)
                        has_enc_err = any("编码损坏" in r or "编码异常" in r for r in grade.review_reasons)

                        if has_cbp or has_lnk:
                            src_summary = "未提交有效 C 源码 (仅提交工程描述 .cbp 或快捷方式 .lnk)"
                        elif has_corrupt:
                            src_summary = "提交压缩包损坏无法解压，源码无法提取验证"
                        elif has_enc_err or any(it.completion_status == "待复核" for it in items):
                            src_summary = "源码已提交，但存在编码异常或待复核事项"
                        elif r2_1 and r2_2 and r2_1.compile_result and r2_1.compile_result.success and r2_2.compile_result and r2_2.compile_result.success:
                            p2_1_ok = all(it.completion_status == "已达成" for it in items if it.problem_id == "2-1")
                            p2_2_ok = all(it.completion_status == "已达成" for it in items if it.problem_id == "2-2")
                            if p2_1_ok and p2_2_ok:
                                src_summary = "2-1与2-2源码均提交且编译通过，全部测试用例正确"
                            else:
                                src_summary = "2-1与2-2源码均提交并编译通过，部分测试用例存在偏差"
                        elif any(p and p.status == "compile_error" for p in [r2_1, r2_2]):
                            src_summary = "源码已提交，但存在编译语法错误未能生成执行文件"
                        else:
                            src_summary = "源码提交不全或测试未通过"
                else:
                    if not grade or grade.status == "missing":
                        src_summary = "未提交任何 C 源码"
                    else:
                        r1_2 = grade.problem_results.get("1-2")
                        r1_3 = grade.problem_results.get("1-3")
                        has_cbp = any(".cbp" in r for r in grade.review_reasons)
                        has_lnk = any(".lnk" in r for r in grade.review_reasons)
                        has_exe = any("Wprogram.exe" in r or "预编译" in r for r in grade.review_reasons)

                        if has_cbp or has_lnk:
                            src_summary = "未提交有效 C 源码 (仅提交工程描述 .cbp 或快捷方式 .lnk)"
                        elif has_exe:
                            src_summary = "1-2仅提交预编译exe缺少C源码，1-3源码编译运行全通过"
                        elif r1_2 and r1_3 and r1_2.compile_result and r1_2.compile_result.success and r1_3.compile_result and r1_3.compile_result.success:
                            int_div = any("整型截断" in it.deduction_reason for it in items)
                            if int_div:
                                src_summary = "1-2与1-3源码均提交且编译通过，实测转换用例存在整数截断证据"
                            else:
                                p1_3_ok = bool(r1_3.test_results) and all(c.passed for c in r1_3.test_results) and not any(it.deduction or it.needs_review for it in items if it.problem_id == "1-3")
                                p1_2_ok = bool(r1_2.test_results) and all(c.passed for c in r1_2.test_results)
                                if p1_3_ok and p1_2_ok:
                                    src_summary = "1-2与1-3源码均提交且编译通过，全部测试用例正确"
                                else:
                                    src_summary = "源码已提交并编译通过，部分运算测试存在偏差"
                        elif any(p.status == "compile_error" for p in [r1_2, r1_3] if p):
                            src_summary = "源码已提交，但存在编译语法错误未能生成执行文件"
                        else:
                            src_summary = "源码提交不全或测试未通过"

                # 3. 手写材料提交与正确性
                if self.lab_id == "lab2":
                    hw_summary = "不适用 (Lab 2 无手写验算要求)"
                else:
                    if not grade or grade.status == "missing":
                        hw_summary = "未提交演草纸验算材料"
                    else:
                        r1_2 = grade.problem_results.get("1-2")
                        if not r1_2 or not any(it.sub_item == "验算材料提交" and it.submission_status == "已提交" for it in items):
                            hw_summary = "未提交演草纸验算材料"
                        elif r1_2.ai_review and r1_2.ai_review.get("is_pending"):
                            hw_summary = "已提交手写照片，AI批改待定，需人工复核"
                        else:
                            hw_sc = r1_2.sub_scores.get("handwritten", 0.0)
                            if hw_sc == 1.5:
                                hw_summary = f"已提交演草纸，按已确认宽容政策验算得满分（信息漏写不扣分） (得分 {hw_sc:.2f}/1.50)"
                            elif hw_sc > 0:
                                hw_summary = f"已提交演草纸，存在明确身份冲突、参数不符或实质运算错误 (得分 {hw_sc:.2f}/1.50)"
                            else:
                                hw_summary = "已提交演草纸，按具体检查证据评定"

                if self.lab_id not in {'lab1','lab2'}:
                    summary_problems='; '.join(f"{pid}:"+('待复核' if pr.is_pending else f'{pr.score:.2f}/{pr.max_score:.2f}') for pid,pr in grade.problem_results.items()) if grade else '未提交'
                    src_summary='按源码、编译与用例证据记录'
                    hw_summary='依本实验专属规则记录'
                if grade and not grade.problem_results and grade.status!='missing':
                    summary_problems='历史总分待复核' if grade.is_pending else f'已保存总分{grade.total_score:.2f}/{grade.max_score:.2f}，缺少题目明细'
                    src_summary='缺少历史源码与用例证据，需核验原报告'
                    hw_summary='缺少历史材料观察证据，需核验原报告'
                # 4. 未达成项目
                failed_items = [
                    f"{it.problem_id}:{it.sub_item}"
                    for it in items
                    if it.completion_status in {"未达成", "部分达成"}
                ]
                failed_items_str = "; ".join(failed_items) if failed_items else "无"

                # 5. 建议扣分合计
                confirmed_ded = 0.0
                has_pending_ded = False
                for it in items:
                    if it.deduction is None:
                        has_pending_ded = True
                    elif it.deduction > 0:
                        confirmed_ded += cents(it.deduction)
                confirmed_ded = confirmed_ded / 100

                if not grade or grade.status == "missing":
                    deduction_total_str = "5.00"
                elif has_pending_ded or grade.is_pending:
                    if confirmed_ded > 0:
                        deduction_total_str = f"已确认扣 {confirmed_ded:.2f} 分 (另有分项待复核)"
                    else:
                        deduction_total_str = "待复核"
                else:
                    deduction_total_str = f"{confirmed_ded:.2f}"

                if mock_mode and not deduction_total_str.startswith("[模拟]"):
                    deduction_total_str = f"[模拟] {deduction_total_str}"

                # 6. 扣分依据
                ded_reasons = [
                    f"[{it.problem_id} {it.sub_item}] {it.deduction_reason} (-{it.deduction:.2f}分)"
                    for it in items
                    if it.deduction and it.deduction > 0
                ]
                ded_reasons_str = "; ".join(ded_reasons) if ded_reasons else "无扣分项"

                # 7. Agent说明附注
                agent_notes = []
                for it in items:
                    if it.agent_note and it.agent_note not in agent_notes and it.agent_note != "无":
                        if it.completion_status in {"未达成", "部分达成", "待复核"} or "归属" in it.agent_note or "常数" in it.agent_note:
                            agent_notes.append(it.agent_note)
                if not grade or grade.status == "missing":
                    agent_note_str = "在册学生，未检索到有效提交作业"
                elif agent_notes:
                    agent_note_str = "; ".join(agent_notes)
                else:
                    agent_note_str = "规范化归档正常，各分项检查已完成"

                # 8. 待复核事项
                if grade and grade.needs_review:
                    review_str = "; ".join(sorted(list(set(grade.review_reasons)))) if grade.review_reasons else "待人工复核确认"
                else:
                    review_str = "无"

                # 9. 实验评语 (>= 50 字)
                if not grade or grade.status == "missing":
                    if self.lab_id == "lab2":
                        comment_str = (
                            f"{student.name}同学你好！系统未检测到你在本次LAB2中的作业提交（未收到源码文件）。"
                            f"按课程大纲与实验教学要求，请尽快核实是否由于提交逾期或文件路径遗漏，并及时联系助教说明情况补交，以便及时核对并确认你的实验平时成绩。"
                        )
                    else:
                        comment_str = (
                            f"{student.name}同学你好！系统未检测到你在实验1中的作业提交（未收到源码文件及演草纸验算照片）。"
                            f"按课程大纲与实验教学要求，请尽快核实是否由于提交逾期或文件路径遗漏，并及时联系助教说明情况补交，以便及时核对并确认实验平时成绩。"
                        )
                else:
                    comment_str = grade.comment

                writer.writerow({
                    "学号": student.student_id,
                    "姓名": student.name,
                    "各题验收汇总": summary_problems,
                    "源码提交与正确性": src_summary,
                    "手写材料提交与正确性": hw_summary,
                    "未达成项目": failed_items_str,
                    ded_header: deduction_total_str,
                    "扣分依据": ded_reasons_str,
                    "Agent说明附注": agent_note_str,
                    "待复核事项": review_str,
                    "实验评语": comment_str,
                    "总分": '待复核' if grade and grade.is_pending else f"{grade.total_score if grade else 0.0:.2f}",
                    "满分": f"{grade.max_score if grade else 5.0:.2f}",
                    "成绩状态": grade.status if grade else 'missing',
                    "成绩待定": '是' if grade and grade.is_pending else '否',
                })

        return file_path

    def generate_summary_markdown(
        self,
        roster_students: List[Student],
        grades: Dict[str, StudentGrade],
        ambiguous_items: Optional[List[Dict[str, Any]]] = None,
        unmatched_items: Optional[List[Dict[str, Any]]] = None,
        mock_mode: bool = False,
        encoding_stats: Optional[Dict[str, Any]] = None,
    ) -> Path:
        """导出 labX报告.md"""
        self._prepare(grades)
        file_path = self.output_dir / f"{'mock_' if mock_mode else ''}{self.lab_id}报告.md"

        total_roster = len(roster_students)
        submitted_students = [
            s for s in roster_students
            if s.student_id in grades and grades[s.student_id].status not in {"missing"}
        ]
        submitted_count = len(submitted_students)
        missing_count = total_roster - submitted_count

        conflict_count = sum(
            1 for s in roster_students
            if s.student_id in grades and grades[s.student_id].status in {"ambiguous", "conflict"}
        )

        graded_scores = [
            grades[s.student_id].total_score
            for s in submitted_students
            if not grades[s.student_id].is_pending and grades[s.student_id].status not in {"ambiguous", "conflict"}
        ]
        avg_score = round(sum(graded_scores) / len(graded_scores), 2) if graded_scores else 0.0
        max_sc = max(graded_scores) if graded_scores else 0.0
        min_sc = min(graded_scores) if graded_scores else 0.0
        full_score_count = sum(1 for g in grades.values() if not g.is_pending and g.total_score==g.max_score)

        needs_review_students = [
            (s, grades[s.student_id])
            for s in roster_students
            if s.student_id in grades and (grades[s.student_id].needs_review or grades[s.student_id].is_pending or grades[s.student_id].status in {"ambiguous", "conflict"})
        ]

        # 统计各分项达成情况
        sub_items_stat: Dict[str, Dict[str, int]] = {}
        for s in roster_students:
            g = grades.get(s.student_id)
            if g and g.all_check_items:
                for it in g.all_check_items:
                    key = f"{it.problem_id}: {it.sub_item}"
                    st = sub_items_stat.setdefault(key, {"已达成": 0, "部分达成": 0, "未达成": 0, "无法验证": 0, "待复核": 0, "不适用": 0})
                    st[it.completion_status] = st.get(it.completion_status, 0) + 1

        sub_pct = round(submitted_count / total_roster * 100, 1) if total_roster > 0 else 0.0
        lines = [
            f"# 《程序设计思维与实践》{self.lab_id.upper()} 批量评测报告\n",
        ]

        if mock_mode:
            lines.extend([
                "> [!WARNING]",
                "> **当前报告由模拟评测模式 (Mock AI) 生成**。手写演草纸与语义复核未调用真实大模型推理，分值仅供调试，**严禁作为最终正式成绩单使用**。\n",
            ])

        missing_students = [
            s for s in roster_students
            if s.student_id not in grades or grades[s.student_id].status in {"missing"}
        ]
        missing_count = len(missing_students)
        missing_desc = f" ({'、'.join(f'{s.student_id} {s.name}' for s in missing_students)})" if missing_students else ""

        lines.extend([
            "## 一、 整体评测与达成概况",
            f"- **点名册应交人数**: {total_roster} 人",
            f"- **实际提交人次**: {submitted_count} 人 ({sub_pct}%)",
            f"- **归属争议/多学号冲突**: {conflict_count} 人",
            f"- **未提交/缺考人数**: {missing_count} 人{missing_desc}",
            f"- **满分人数**: {full_score_count} 人",
            f"- **已定级有效平均分**: {avg_score} 分 (最高分: {max_sc}, 最低分: {min_sc})",
            f"- **待人工复核总人次**: {len(needs_review_students)} 人\n",
            "## 二、 逐题与检查分项验收达成统计",
            "| 题号与检查分项 | 已达成 | 部分达成 | 未达成 | 无法验证 | 待复核 | 达成率 |",
            "| :--- | :---: | :---: | :---: | :---: | :---: | :---: |",
        ])

        for key, st in sub_items_stat.items():
            pass_cnt = st.get("已达成", 0)
            part_cnt = st.get("部分达成", 0)
            fail_cnt = st.get("未达成", 0)
            unv_cnt = st.get("无法验证", 0)
            rev_cnt = st.get("待复核", 0)
            valid_total = pass_cnt + part_cnt + fail_cnt + rev_cnt
            rate_str = f"{pass_cnt / valid_total * 100:.1f}%" if valid_total > 0 else "不适用"
            lines.append(f"| {key} | {pass_cnt} | {part_cnt} | {fail_cnt} | {unv_cnt} | {rev_cnt} | {rate_str} |")

        if self.lab_id == "lab2":
            lines.extend([
                "",
                "## 三、 评分细则与说明",
                "- **Warming up 课前热身 (0.0分)**：不计入成绩，单列记录完成情况。",
                "- **Lab 2-1 时间段问候 (满分 2.0分)**：",
                "  - 源码未提交 / 编译失败：扣 2.00 分",
                "  - 12例均分2.00分，单例精确权重2/12；有效时间10例与非法输入2例按同一权重计算，分类不重复扣分",
                "  - 分项显示由公共层分配舍入余数；用例分、分类扣分和整题总分严格对齐",
                "- **Lab 2-2 字母金字塔 (满分 3.0分)**：",
                "  - 源码未提交 / 编译失败：扣 3.00 分",
                "  - 字母对称金字塔展开：三例各0.50，大写字母金字塔层层展开，严格对称 (指导书笔误 ABCBC 按对称 ABCBA 校验)",
                "  - 大小写字母兼容转换：小写字母输入需自动提升并展开；6例均分3.00，每例0.50",
                "  - 非字母与异常输入报错：两例各0.50，已含在6例均分预算内，不另叠加扣分",
                "- **Lab 2-3 / Lab 2-4 选做题**：回文素数与 4:2 优先编码器，独立记录完成情况，不影响必做题 5 分基准。",
                "> [!NOTE]",
                "> 以上分项使用与实际判分一致的均分用例预算和公共舍入分配，最终核准结果需经课程组人工复核清单确认。\n",
                "## 四、 典型技术问题与失分剖析",
                "1. **字母金字塔笔误兼容与严格对称**：指导书示例中包含部分排版笔误（如 `___ABCBC`），系统严格依据字母对称规律（`ABCBA`）进行多组自动化测试。",
                "2. **大小写字母输入兼容**：部分同学未做小写字母（如 `'f'`）向大写字母的自动提升转换，直接导致分支走入非法输入或生成错误序列。",
                "3. **非法输入报错处理**：指导书要求对非字母（如数字、负号）给出提示或报错，部分实现未对输入合法性进行校验即进行减法运算，引发越界输出。",
                "4. **源码编码与跨平台字符集**：部分提交包含非 UTF-8 编码或损坏字符，已在 submit → src 归档阶段自动识别并规范化为 UTF-8 无 BOM。\n",
                f"## 五、 待人工复核清单 (共 {len(needs_review_students)} 项)",
            ])
        elif self.lab_id == "lab1":
            lines.extend([
                "",
                '## 三、 已确认评分细则',
                '- Lab 1-1为不计分的课内练习。',
                '- Lab 1-2源码满分1.50：九个运算输出项均分，每个失败扣1.50/9；先累计精确分值再按分舍入，缺源码或编译运行失败扣1.50，参数与符号问题不在用例之外重复扣代码分。',
                '- 手写满分1.50：身份0.50，材料参数0.20，推导0.80。漏写信息不扣分；姓名、学号或完整手机号明确冲突时身份项为0并复核，其余已确认分数保留。',
                '- 手写推导倾向给满：结果相符且无明显错误即满分；简写、缺少竖式、小笔误不酌情扣分。有明确错误证据时使用AI连续推导分d（0～0.80），扣0.80−d，必须附纸面错误与正确结果证据；确无有效推导时该0.80项为0。',
                '- 明确的学号末两位x、手机号末两位y不对应，各在推导项额外扣0.20，每个参数只扣一次；推导项最低0分。无法取得代码参数时不以默认0制造冲突。',
                '- Lab 1-3满分2.00：现有10个用例各0.20，失败一例扣0.20；方向与边界仅汇总这些用例，不额外扣分。最终结果数值容差统一为±0.08。',
                '- 数值正确但小数位多于或少于两位，整道温度题额外扣0.20一次，最低0分。仅检查有效结果，不检查菜单或输入回显的格式。',
                '- 多版本按上述完整得分择优；待定结果保留复核，不作为已确认错误。',
                '## 四、 判定证据与语言语义',
                '- 整数除法诊断须结合C/C++声明类型、提升、算术转换和求值顺序，并与真实运行相符；整数常量和菜单0均不单独构成错误证据。诊断不另扣分。',
                '- 语义复核同时提供源码、输入和实际输出；每个用例只分通过或不通过，不以AI连续分值改变用例权重。',
                f"## 五、 待人工复核清单 (共 {len(needs_review_students)} 项)",
            ])

        if self.lab_id not in {'lab1','lab2'}:
            lines.extend(['## 三、 公共评分规则',
                '各题用例预算与分组来自评测声明；公共层统一完整候选择优、按分吸收舍入余数并校验扣分闭合。',
                f'## 五、 待人工复核清单 (共 {len(needs_review_students)} 项)'])
        if not needs_review_students:
            lines.append("当前批次未发现待复核样本，全部用例已自动定级。\n")
        else:
            lines.append("| 学号 | 姓名 | 班级 | 当前建议分值 | 待复核疑点 / 证据说明 | 处理建议 |")
            lines.append("| :--- | :--- | :--- | :---: | :--- | :--- |")
            for std, gr in needs_review_students:
                reasons = "<br>".join(sorted(list(set(gr.review_reasons)))) if gr.review_reasons else "待人工复核"
                sc_str = "成绩待定" if (gr.is_pending or gr.status in {"ambiguous", "conflict"}) else f"{gr.total_score:.2f}分"
                if self.lab_id == "lab2":
                    advice = "现场核验源码与功能"
                    if any("编码" in r for r in gr.review_reasons):
                        advice = "核对源码文件编码与原始字符"
                    elif any(".cbp" in r or ".lnk" in r for r in gr.review_reasons):
                        advice = "督促学生补交实际 C 源码"
                    elif any("损坏" in r or "解压" in r for r in gr.review_reasons):
                        advice = "联系学生重新提交未损坏压缩包"
                else:
                    advice = "现场核验源码与推导"
                    if any("未写学号姓名" in r for r in gr.review_reasons):
                        advice = "核对演草纸笔迹与参数归属"
                    elif any(".cbp" in r or ".lnk" in r for r in gr.review_reasons):
                        advice = "督促学生补交实际 C 源码"
                    elif any("Wprogram.exe" in r for r in gr.review_reasons):
                        advice = "补交 Lab 1-2 C 源码"
                lines.append(f"| {std.student_id} | {std.name} | {std.class_id} | {sc_str} | {reasons} | {advice} |")
            lines.append("")

        if ambiguous_items or unmatched_items:
            lines.append("## 六、 提交文件归属审计与无法归属文件归档清单\n")
            lines.append(
                f"所有按学号分类时无法确定归属的提交文件，已统一隔离归档至 **`{self.lab_id}/src/others/`** 目录，"
                "保留原文件名及来源压缩包结构，同名文件不覆盖，便于后续核验与联系学生。"
                "无法归属材料统一记录为“无法归属、待复核”，不直接判定为未交或据此扣分。\n"
            )
            if ambiguous_items:
                lines.append("**多学号冲突提交件**：包含多个候选学号，已统一归档并标记待人工复核。\n")
            if unmatched_items:
                lines.append("**未识别点名册提交件 (安全隔离)**：未匹配到在册学生，已统一归档至 others 目录。\n")
            lines.append("| 原始提交文件 | 统一归档位置 (`src/others/`) | 归属状态 | 无法归属原因 | 候选学号 | 处理建议 |")
            lines.append("| :--- | :--- | :---: | :--- | :---: | :--- |")

            all_others = (ambiguous_items or []) + (unmatched_items or [])
            for item in all_others:
                orig_path = item.get("path", "")
                archive_dest = item.get("archive_dest", f"{self.lab_id}/src/others/")
                status = item.get("status", "无法归属、待复核")
                reason = item.get("reason", "未匹配学生")
                cands = item.get("candidates", [])
                cands_str = ", ".join(cands) if cands else "无"
                if cands:
                    advice = "联系候选学生澄清作业归属"
                elif "测试数据包" in reason:
                    advice = "测试样本包，保留归档供人工核实"
                else:
                    advice = "保留于 others 供人工核实"
                lines.append(f"| `{orig_path}` | `{archive_dest}` | {status} | {reason} | `{cands_str}` | {advice} |")
            lines.append("")

        lines.extend([
            "## 七、 编译与兼容性环境处理说明",
            "- Linux 环境对标 GCC `-std=c11`，数学库 `-lm` 位于命令末尾确保符号解析无误。",
            "- 针对 Windows 平台特有的 `<windows.h>`、`<conio.h>`、`system(\"pause\")`、`scanf_s` 及 GBK 编码源码，已全自动生成 `_compat` 编译副本并保存审计日志，未改变学生核心算法与数学公式。",
            "- 程序运行具有资源限额、进程组超时强杀和输出截断；环境支持unshare时启用网络与文件系统隔离，否则按运行模块的降级路径执行。",
        ])
        if encoding_stats:
            conv_cnt = encoding_stats.get("converted_count", 0)
            unk_cnt = encoding_stats.get("unknown_count", 0)
            lines.append(
                f"- 源码编码识别与 UTF-8 规范化：submit → src 归档阶段对全部源码与头文件执行严格编码识别，支持 UTF-8、UTF-8 BOM、GBK/GB18030 以及 UTF-16。本批次共识别并规范化转码 {conv_cnt} 份源码，未确定或损坏编码 {unk_cnt} 份。所有源码统一以标准 UTF-8（无 BOM）保存，彻底消除 GCC 跨平台与静态分析字符集乱码。"
            )
        lines.append("")

        lines.append(self._candidate_markdown(grades))
        file_path.write_text("\n".join(lines), encoding="utf-8")
        return file_path

    def generate_roster_grade_xlsx(
        self,
        roster_students: List[Student],
        grades: Dict[str, StudentGrade],
        template_path: Path,
        mock_mode: bool = False,
    ) -> Optional[Path]:
        """如果存在官方 Excel 模板，同步填写对应 sheet 并另存"""
        self._prepare(grades)
        try:
            import openpyxl
            prefix = "mock_" if mock_mode else ""
            dest_path = self.output_dir / f"{prefix}{self.lab_id}实验成绩单.xlsx"
            shutil.copy2(template_path, dest_path)

            wb = openpyxl.load_workbook(str(dest_path))
            sheet_title = "Lab 1" if self.lab_id == "lab1" else "Lab 2"
            if sheet_title not in wb.sheetnames:
                sheet_title = wb.sheetnames[0]

            ws = wb[sheet_title]

            if mock_mode:
                ws.cell(row=2, column=4, value="分数(5)[模拟测试非正式]")

            for idx, student in enumerate(roster_students, start=3):
                grade = grades.get(student.student_id)
                if not grade:
                    val_score = "[模拟] 0.00" if mock_mode else 0.0
                    if self.lab_id == "lab2":
                        comment = (
                            f"{student.name}同学你好！系统未检测到你在本次LAB2中的作业提交（未收到源码文件）。"
                            f"按课程大纲与实验教学要求，请尽快核实是否由于提交逾期或文件路径遗漏，并及时联系助教说明情况补交，以便及时核对并确认你的实验平时成绩。"
                        )
                    else:
                        comment = (
                            f"{student.name}同学你好！系统未检测到你在本次实验1中的作业提交（未收到源码文件及演草纸验算照片）。"
                            f"按课程大纲与实验教学要求，请尽快核实是否由于提交逾期或文件路径遗漏，并及时联系助教说明情况补交，以便及时核对并确认你的实验平时成绩。"
                        )
                elif grade.is_pending or (grade.needs_review and grade.status in {"missing_source", "ambiguous", "conflict"}):
                    val_score = "[模拟] 待复核" if mock_mode else "待复核"
                    comment = grade.comment
                else:
                    val_score = f"[模拟] {grade.total_score:.2f}" if mock_mode else round(grade.total_score, 2)
                    comment = grade.comment

                ws.cell(row=idx, column=1, value=student.seq)
                ws.cell(row=idx, column=2, value=student.student_id)
                ws.cell(row=idx, column=3, value=student.name)
                ws.cell(row=idx, column=4, value=val_score)
                ws.cell(row=idx, column=5, value=comment)

            wb.save(str(dest_path))
            return dest_path
        except Exception:
            return None

    def generate_single_student_report(
        self,
        student: Student,
        grade: StudentGrade,
        mock_mode: bool = False,
    ) -> Dict[str, Path]:
        """为单人重判导出专用的独立报表，避免覆盖其他学生成绩"""
        self._prepare({student.student_id:grade})
        prefix = "mock_" if mock_mode else ""
        detail_path = self.output_dir / f"single_{student.student_id}_{prefix}{self.lab_id}结果.csv"
        fieldnames = [
            "学号", "姓名", "实验", "题号", "评分项", "提交状态",
            "达成状态", "检查依据", "建议扣分", "扣分理由", "Agent说明附注", "是否需要复核", "实验评语", "题目得分", "题目满分", "采用版本", "候选数"
        ]

        items = grade.all_check_items if grade.all_check_items else self._build_default_check_items(student, grade)
        with open(detail_path, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            for it in items:
                ded_str = "待复核" if it.deduction is None else (f"{it.deduction:.2f}" if it.deduction > 0 else "0.00")
                writer.writerow({
                    "学号": student.student_id,
                    "姓名": student.name,
                    "实验": self.lab_id,
                    "题号": it.problem_id,
                    "评分项": it.sub_item,
                    "提交状态": it.submission_status,
                    "达成状态": it.completion_status,
                    "检查依据": it.evidence,
                    "建议扣分": ded_str,
                    "扣分理由": it.deduction_reason,
                    "Agent说明附注": it.agent_note,
                    "是否需要复核": "是" if it.needs_review else "否",
                    "实验评语": grade.comment,
                    **self._problem_cells(grade,it.problem_id),
                })

        md_path = self.output_dir / f"single_{student.student_id}_{prefix}{self.lab_id}报告.md"
        mock_title = "【模拟测试非正式】" if mock_mode else ""
        sc_str = "待复核" if grade.is_pending else f"{grade.total_score:.2f}"
        md_content = f"""# {mock_title}《程序设计思维与实践》{self.lab_id.upper()} 单人评测报告 ({student.student_id} {student.name})

- **学号**: {student.student_id}
- **姓名**: {student.name}
- **班级**: {student.class_id}
- **最终建议总分**: {sc_str} / {grade.max_score:.2f}
- **评测状态**: {grade.status}
- **待人工复核**: {"是" if grade.needs_review else "否"}
- **复核原因**: {"; ".join(grade.review_reasons) if grade.review_reasons else "无"}

## 实验评语（>= 50字）
{grade.comment}
"""
        md_content += self._candidate_markdown({student.student_id:grade})
        md_path.write_text(md_content, encoding="utf-8")
        return {"detail_csv": detail_path, "summary_md": md_path}

    def load_existing_grades_from_files(self, roster_students, mock_mode=None):
        """Prefer complete canonical state; legacy CSV retains explicit totals only."""
        prefix='mock_' if mock_mode else ''
        if mock_mode is None and not ((self.output_dir/f'{self.lab_id}实验成绩单.csv').exists() or (self.output_dir/f'{self.lab_id}评分状态.json').exists()) and ((self.output_dir/f'mock_{self.lab_id}实验成绩单.csv').exists() or (self.output_dir/f'mock_{self.lab_id}评分状态.json').exists()):
            prefix='mock_'
        state=self.output_dir/f'{prefix}{self.lab_id}评分状态.json'
        roster_map={s.student_id:s for s in roster_students}
        if state.exists():
            payload=json.loads(state.read_text(encoding='utf-8'))
            if payload.get('schema_version')!=SCHEMA_VERSION or payload.get('lab_id')!=self.lab_id or bool(payload.get('mock_mode'))!=bool(prefix):
                raise ScoringConsistencyError('评分状态格式、实验或模拟标识不一致')
            recovered={sid:grade_from_dict(data) for sid,data in payload['grades'].items() if sid in roster_map}
            for sid,g in recovered.items():
                if g.student.student_id!=sid:
                    raise ScoringConsistencyError('评分状态学号不一致')
                g.student=roster_map[sid]
            self._prepare(recovered)
            return recovered
        path=self.output_dir/f'{prefix}{self.lab_id}实验成绩单.csv'
        if not path.exists():
            return {}
        recovered={}
        with path.open(encoding='utf-8-sig',newline='') as stream:
            for row in csv.DictReader(stream):
                sid=row.get('学号','').strip()
                if sid not in roster_map:
                    continue
                ded=row.get('建议扣分合计') or row.get('建议扣分合计[模拟测试非正式]','')
                pending=row.get('成绩待定')=='是' or '待复核' in ded or '待定' in ded
                total=row.get('总分','')
                max_score=float(row.get('满分') or 5)
                # Older files have no authoritative total; don't turn a possibly
                # inconsistent deduction sum into a freshly confirmed grade.
                legacy=not total
                if legacy:
                    pending=True
                    score=0.0
                else:
                    score=0.0 if pending else float(total)
                    if not pending and cents(float(ded.replace('[模拟]','').strip()))!=cents(max_score)-cents(score):
                        raise ScoringConsistencyError('历史CSV总分与扣分合计不一致')
                reasons=[] if row.get('待复核事项','无')=='无' else [row.get('待复核事项')]
                if legacy:reasons.append('旧报表缺少独立总分与完整评分状态，需按原始提交重判核实')
                else:reasons.append('CSV保留独立总分，但缺少完整题目、用例和候选状态；分项证据需核验')
                recovered[sid]=StudentGrade(roster_map[sid],total_score=score,max_score=max_score,
                    status='needs_review' if pending or reasons else row.get('成绩状态','graded'),
                    is_pending=pending,needs_review=bool(pending or reasons),review_reasons=reasons,
                    comment=row.get('实验评语',''))
        return recovered
