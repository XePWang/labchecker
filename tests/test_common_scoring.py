"""Cross-lab scoring contracts and real archive-to-report round trips."""
import contextlib
import copy
import csv
import io
import itertools
import json
import tempfile
import unittest
from dataclasses import asdict
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from unittest.mock import patch

from checker.config import WORK_DIR
from checker.labs.base import BaseLabJudge
from checker.labs.lab1 import Lab1Judge
from checker.labs.lab2 import Lab2Judge
from checker.models import CheckItem,CompileResult,ProblemResult,Student,StudentGrade,TestCaseResult
from checker.reporter import LabReporter
from checker.scoring import ScoringConsistencyError,finalize_problem,finalize_grade,scored_problem

class ThirdLab(BaseLabJudge):
    """A fresh lab declares facts only; no local rounding or candidate max()."""
    def __init__(self): super().__init__('lab99',1.0)
    def collect_compile_tasks(self,*args): return []
    @scored_problem('99-1')
    def evaluate(self,student,sources,build_dir,precompiled=None):
        facts=json.loads(sources[0].read_text())
        results=[TestCaseResult(str(i),'','',str(ok),ok,0,0) for i,ok in enumerate(facts['passed'])]
        p=ProblemResult('99-1',0,1,compile_result=CompileResult(True),test_results=results,
            check_items=[CheckItem('99-1','用例','已提交','部分达成','实际检查事实',0),
                         CheckItem('99-1','格式','已提交','已达成','独立格式事实',facts.get('penalty',0))],
            case_budget=1.0,case_groups={'用例':list(range(len(results)))})
        if facts.get('pending'):
            p.is_pending=True;p.check_items[0].deduction=None
        return p
    def _grade_student(self,student,files,build_dir,mock_ai=False,precompiled=None):
        p=self.evaluate(student,[f.path for f in files],build_dir,precompiled)
        return StudentGrade(student,max_score=1,problem_results={'99-1':p},status='graded')
    def generate_comment(self,g): return '受控第三实验仅声明用例预算和事实，完整版本评定、候选择优、舍入余数分配和报表一致性全部使用公共实现。'

class CommonScoringTests(unittest.TestCase):
    def setUp(self):
        WORK_DIR.mkdir(exist_ok=True,parents=True)
        self.tmp=tempfile.TemporaryDirectory(prefix='common_score_',dir=WORK_DIR)
        self.root=Path(self.tmp.name)
        self.addCleanup(self.tmp.cleanup)
        self.student=Student(1,'2026999901','一致性测试甲')

    def write(self,name,text):
        p=self.root/name;p.parent.mkdir(parents=True,exist_ok=True);p.write_text(text,encoding='utf-8');return p

    def closed(self,p):
        self.assertFalse(p.is_pending)
        self.assertEqual(sum(Decimal(str(i.deduction)) for i in p.check_items),Decimal(str(p.max_score))-Decimal(str(p.score)))

    def test_all_pass_fail_patterns_close_for_multiple_case_counts(self):
        # All 3/6/8/9/12-case pass/fail patterns, including half-cent boundaries.
        for n,budget in [(3,1),(6,3),(8,1),(9,1.5),(12,2)]:
            groups={'前组':list(range(n//2)),'后组':list(range(n//2,n))}
            for mask in range(1<<n):
                flags=[bool(mask&(1<<i)) for i in range(n)]
                cases=[TestCaseResult(str(i),'','','',ok,0,0) for i,ok in enumerate(flags)]
                p=ProblemResult('generic',0,budget,test_results=cases,case_budget=budget,case_groups=groups,
                    check_items=[CheckItem('generic',name,'已提交','部分达成','受控事实',0) for name in groups])
                finalize_problem(p)
                expected=(Decimal(str(budget))*sum(flags)/n).quantize(Decimal('.01'),rounding=ROUND_HALF_UP)
                self.assertEqual(Decimal(str(p.score)),expected,(n,mask))
                self.closed(p)
                self.assertEqual(sum(Decimal(str(c.score)) for c in cases),expected)
                self.assertEqual(sum(Decimal(str(c.max_score)) for c in cases),Decimal(str(budget)))
                self.assertTrue(all(0<=c.score<=c.max_score for c in cases))
                if mask in (0,1,(1<<n)-1):
                    before=asdict(p);finalize_problem(p);self.assertEqual(before,asdict(p))

    def test_rejects_double_counted_or_omitted_case_membership(self):
        p=ProblemResult('x',0,1,case_budget=1,case_groups={'a':[0,0]},
            test_results=[TestCaseResult('c','','','',False,0,0)],check_items=[CheckItem('x','a','已提交','未达成','事实',0)])
        with self.assertRaises(ScoringConsistencyError):finalize_problem(p)

    def test_substantive_score_conflict_is_not_absorbed_as_rounding(self):
        p=ProblemResult('x',.7,1,check_items=[CheckItem('x','明确错误','已提交','未达成','事实',.5)])
        with self.assertRaises(ScoringConsistencyError):finalize_problem(p)
        reporter=LabReporter('lab99',self.root/'bad_report')
        g=StudentGrade(self.student,max_score=1,problem_results={'x':p},status='graded')
        with self.assertRaises(ScoringConsistencyError):reporter.generate_all_reports([self.student],{self.student.student_id:g})
        self.assertEqual(list(reporter.output_dir.iterdir()),[])

    def test_new_lab_inherits_finalization_and_final_score_selection(self):
        from checker.models import SubmissionFile
        j=ThirdLab()
        a=self.write('a.json',json.dumps({'passed':[True]*10,'penalty':.2}))
        b=self.write('b.json',json.dumps({'passed':[True]*9+[False]}))
        for paths in [[a,b],[b,a]]:
            p=j.evaluate(self.student,paths,self.root/'build')
            self.assertEqual(p.score,.9)
            self.assertEqual(p.chosen_candidate['src_path'],str(b))
            self.assertEqual(len(p.candidate_evaluations),2)
            self.assertEqual(sorted(c['score'] for c in p.candidate_evaluations),[.8,.9])
            self.assertNotIn('0.00/',p.chosen_candidate['reason'])
            self.closed(p)
        g=j.grade_student(self.student,[SubmissionFile(b,b.name,'other')],self.root/'build')
        self.assertEqual(g.total_score,.9);self.assertGreaterEqual(len(g.comment),50)
        report=LabReporter('lab99',self.root/'generic_report')
        files=report.generate_all_reports([self.student],{self.student.student_id:g})
        text=files['summary_md'].read_text()
        self.assertIn('公共评分规则',text);self.assertNotIn('Lab 1-2源码',text)

    def test_pending_candidates_never_poison_a_confirmed_full_score(self):
        j=ThirdLab()
        good=self.write('good.json',json.dumps({'passed':[True]*10}))
        pending=self.write('pending.json',json.dumps({'passed':[True]*9+[False],'pending':True}))
        partial=self.write('partial.json',json.dumps({'passed':[True]*8+[False]*2}))
        for sources in [[pending,good],[good,pending]]:
            p=j.evaluate(self.student,sources,self.root/'build')
            self.assertEqual(p.score,1);self.assertFalse(p.is_pending)
            self.assertEqual(p.chosen_candidate['src_path'],str(good))
            self.assertTrue(any(c['is_pending'] for c in p.candidate_evaluations))
        p=j.evaluate(self.student,[pending,partial],self.root/'build')
        self.assertTrue(p.is_pending);self.assertEqual(p.score,.8)
        g=StudentGrade(self.student,max_score=1,status='graded',problem_results={'99-1':p})
        finalize_grade(g)
        report=LabReporter('lab99',self.root/'pending_report')
        files=report.generate_all_reports([self.student],{self.student.student_id:g})
        rows=list(csv.DictReader(files['roster_csv'].open(encoding='utf-8-sig')))
        self.assertEqual(rows[0]['总分'],'待复核')
        restored=report.load_existing_grades_from_files([self.student])[self.student.student_id]
        self.assertTrue(restored.is_pending)
        self.assertEqual(restored.problem_results['99-1'].score,.8)

    @staticmethod
    def greet_code(condition='0'):
        return '''#include <stdio.h>
int main(){int h; if(scanf("%d",&h)!=1)return 1;
if(CONDITION){puts("unhandled");return 0;}
if(h<0||h>23)puts("Invalid input");else if(h<12)puts("Good morning");else if(h==12)puts("Good noon");else if(h<18)puts("Good afternoon");else if(h<21)puts("Good evening");else puts("Good night");return 0;}
'''.replace('CONDITION',condition)

    @staticmethod
    def pyramid_code(condition='0'):
        return r'''#include <stdio.h>
#include <ctype.h>
int main(){char c;if(scanf(" %c",&c)!=1)return 1;
if(CONDITION){puts("unhandled");return 0;}
c=toupper((unsigned char)c);if(c<'A'||c>'Z'){puts("Invalid character!");return 0;}
for(char row='A';row<=c;row++){for(char k='A';k<=row;k++)putchar(k);for(char k=row-1;k>='A';k--)putchar(k);putchar('\n');}return 0;}
'''.replace('CONDITION',condition)

    def test_lab2_real_program_scores_and_deductions_use_one_weight_plan(self):
        j=Lab2Judge()
        cases=[('2-1',self.greet_code('h<0||h>23'),1.67,.33),
               ('2-1',self.greet_code('h==8'),1.83,.17),
               ('2-1',self.greet_code('h==8||h==11||h==13'),1.5,.5),
               ('2-2',self.pyramid_code("c=='C'"),2.5,.5),
               ('2-2',self.pyramid_code("c=='-'"),2.5,.5)]
        for index,(pid,source,score,ded) in enumerate(cases):
            with self.subTest(pid=pid,index=index):
                src=self.write(f'case_{index}.c',source)
                fn=j._grade_lab2_1 if pid=='2-1' else j._grade_lab2_2
                p=fn(self.student,[src],self.root/'build',mock_ai=False)
                self.assertEqual(p.score,score);self.closed(p)
                self.assertEqual(sum(Decimal(str(c.score)) for c in p.test_results),Decimal(str(score)))
                self.assertEqual(sum(Decimal(str(i.deduction)) for i in p.check_items),Decimal(str(ded)))

    def test_optional_versions_all_evaluated_and_selected_by_completion(self):
        j=Lab2Judge()
        full=self.write('encoder_good.c','#include <stdio.h>\nint main(){char s[10];scanf("%9s",s);if(s[0]==\'1\')puts("11");else puts("inactive");}')
        partial=self.write('encoder_partial.c','#include <stdio.h>\nint main(){puts("11");}')
        broken=self.write('encoder_bad.c','int main(){missing;}')
        for sources in [[full,partial,broken],[broken,partial,full]]:
            p=j._grade_lab2_4(self.student,sources,self.root/'optional')
            self.assertEqual(p.score,0);self.assertEqual(p.max_score,0)
            self.assertEqual(len(p.candidate_evaluations),3)
            self.assertEqual(p.chosen_candidate['src_path'],str(full))
            self.assertEqual(p.chosen_candidate['passed_cases'],2)

    def test_archive_cli_all_reports_and_single_student_merge_preserve_other_students(self):
        import openpyxl
        from checker import cli
        from checker.models import SubmissionFile
        # Exercise real normalization, compilation, grading and the public CLI export path.
        roster=self.root/'roster.xlsx';wb=openpyxl.Workbook();ws=wb.active
        ws.append(['序号','学号','姓名']);ws.append([1,self.student.student_id,self.student.name])
        other=Student(2,'2026999902','一致性测试乙');ws.append([2,other.student_id,other.name]);wb.save(roster)
        template=self.root/'template.xlsx';wb=openpyxl.Workbook();wb.active.title='Lab 2';wb.save(template)
        submit=self.root/'lab2/submit';submit.mkdir(parents=True)
        for sid in [self.student.student_id,other.student_id]:
            self.write(f'lab2/submit/{sid}_Lab2-1.c',self.greet_code('h<0||h>23'))
            self.write(f'lab2/submit/{sid}_Lab2-2.c',self.pyramid_code())
        with patch.object(cli,'PROJECT_ROOT',self.root),contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(cli.run_lab('lab2',roster,template,compile_workers=1),0)
        output=self.root/'lab2/output';report=LabReporter('lab2',output)
        people=[self.student,other];before=report.load_existing_grades_from_files(people)
        self.assertEqual(before[other.student_id].total_score,4.67)
        saved_other=asdict(before[other.student_id])
        roster_rows=list(csv.DictReader((output/'lab2实验成绩单.csv').open(encoding='utf-8-sig')))
        self.assertEqual(roster_rows[0]['总分'],'4.67');self.assertEqual(roster_rows[0]['建议扣分合计'],'0.33')
        details=list(csv.DictReader((output/'lab2结果.csv').open(encoding='utf-8-sig')))
        self.assertEqual(sum(Decimal(r['建议扣分']) for r in details if r['学号']==other.student_id),Decimal('.33'))
        xlsx=openpyxl.load_workbook(output/'lab2实验成绩单.xlsx',data_only=True)
        self.assertEqual(xlsx.active.cell(4,4).value,4.67)
        text=(output/'lab2报告.md').read_text()
        self.assertIn('12例均分2.00',text);self.assertIn('候选版本与最终选择',text)
        self.assertNotIn('时间段问候逻辑 (1.50分)',text)
        # Change only one student's code and regrade. The other student's entire
        # result, candidate audit, checks, score and comment must remain intact.
        self.write(f'lab2/submit/{self.student.student_id}_Lab2-1.c',self.greet_code())
        with patch.object(cli,'PROJECT_ROOT',self.root),contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(cli.run_lab('lab2',roster,template,single_student_id=self.student.student_id,compile_workers=4),0)
        after=report.load_existing_grades_from_files(people)
        self.assertEqual(after[self.student.student_id].total_score,5)
        self.assertEqual(asdict(after[other.student_id]),saved_other)
        self.assertEqual(len(list(csv.DictReader((output/'lab2结果.csv').open(encoding='utf-8-sig')))),20)

    def test_corrupt_state_does_not_silently_fall_back_to_old_csv(self):
        reporter=LabReporter('lab99',self.root/'corrupt_state')
        (reporter.output_dir/'lab99评分状态.json').write_text('{invalid')
        with self.assertRaises(json.JSONDecodeError):reporter.load_existing_grades_from_files([self.student])

    def test_legacy_deduction_only_csv_is_pending_instead_of_recomputed_score(self):
        reporter=LabReporter('lab2',self.root/'legacy')
        p=reporter.output_dir/'lab2实验成绩单.csv'
        with p.open('w',encoding='utf-8-sig',newline='') as f:
            w=csv.DictWriter(f,fieldnames=['学号','建议扣分合计','实验评语']);w.writeheader();w.writerow({'学号':self.student.student_id,'建议扣分合计':'.50'})
        restored=reporter.load_existing_grades_from_files([self.student])[self.student.student_id]
        self.assertTrue(restored.is_pending)
        self.assertTrue(any('旧报表' in r for r in restored.review_reasons))

    def test_valid_json_with_contradictory_totals_is_rejected(self):
        j=ThirdLab();src=self.write('valid.json',json.dumps({'passed':[True]*9+[False]}))
        p=j.evaluate(self.student,[src],self.root/'build')
        g=StudentGrade(self.student,total_score=.9,max_score=1,status='graded',problem_results={'99-1':p},comment=j.generate_comment(None))
        reporter=LabReporter('lab99',self.root/'tampered')
        state=reporter.generate_all_reports([self.student],{self.student.student_id:g})['state_json']
        data=json.loads(state.read_text());data['grades'][self.student.student_id]['total_score']=.7
        state.write_text(json.dumps(data))
        with self.assertRaises(ScoringConsistencyError):reporter.load_existing_grades_from_files([self.student])

    def test_lab1_cli_versions_handwriting_and_report_evidence_agree(self):
        import openpyxl
        from checker import cli
        roster=self.root/'roster1.xlsx';wb=openpyxl.Workbook();ws=wb.active
        ws.append(['序号','学号','姓名']);ws.append([1,self.student.student_id,self.student.name]);wb.save(roster)
        template=self.root/'template1.xlsx';wb=openpyxl.Workbook();wb.active.title='Lab 1';wb.save(template)
        self.write(f'lab1/submit/{self.student.student_id}_Lab1-2.c',
            '#include <stdio.h>\nint main(){short x=1,y=88;printf("%d %d %d %u %d %d %d %d %d\\n",!x,x||y,x&&y,(unsigned short)~x,x|y,x&y,~x,x|y,x&y);}')
        # Two independent versions: the correctly formatted version must win.
        for suffix,fmt in [('more','%.3f'),('right','%.2f')]:
            self.write(f'lab1/submit/{self.student.student_id}_Lab1-3_{suffix}.c',
                '#include <stdio.h>\nint main(){int m;double t;scanf("%d%lf",&m,&t);'
                f'printf("{fmt}\\n",m?9*t/5+32:5*(t-32)/9);}}')
        self.write(f'lab1/submit/{self.student.student_id}_handwritten.png','placeholder image'*30)
        observation=dict(task_id='lab1_2_handwritten',student_id_observed='2026999999',name_observed='',
            phone_observed='88',x_observed=1,y_observed=88,info_score=.5,xy_match_score=.2,
            derivation_score=.8,suggested_score=1.5,needs_review=False,reasoning='手写观察受控事实',
            clear_error_operators=[],error_evidence=[],no_valid_derivation=False,derivation_uncertain=False)
        with patch.object(cli,'PROJECT_ROOT',self.root),patch('checker.labs.lab1.HandwrittenReviewer.review_images',return_value=observation),contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(cli.run_lab('lab1',roster,template,mock_ai=True,compile_workers=4),0)
        reporter=LabReporter('lab1',self.root/'lab1/output')
        restored=reporter.load_existing_grades_from_files([self.student],mock_mode=True)[self.student.student_id]
        self.assertEqual(restored.total_score,4.5);self.assertTrue(restored.needs_review);self.assertFalse(restored.is_pending)
        p=restored.problem_results['1-3'];self.assertEqual(p.score,2)
        self.assertEqual(sorted(c['score'] for c in p.candidate_evaluations),[1.8,2])
        self.assertIn('right',p.chosen_candidate['filename'])
        for problem in restored.problem_results.values():self.closed(problem)
        files=reporter.output_dir
        rows=list(csv.DictReader((files/'mock_lab1实验成绩单.csv').open(encoding='utf-8-sig')))
        self.assertEqual(rows[0]['总分'],'4.50')
        self.assertIn('全部测试用例正确',rows[0]['源码提交与正确性'])
        self.assertIn('明确身份冲突',rows[0]['手写材料提交与正确性'])
        wb=openpyxl.load_workbook(files/'mock_lab1实验成绩单.xlsx',data_only=True)
        self.assertEqual(wb.active.cell(3,4).value,'[模拟] 4.50')
        text=(files/'mock_lab1报告.md').read_text()
        self.assertIn('容差统一为±0.08',text);self.assertIn(p.chosen_candidate['filename'],text)
        # Report entrypoints cannot silently repair a tampered final score.
        restored.total_score=4.6
        with self.assertRaises(ScoringConsistencyError):reporter.generate_all_reports([self.student],{self.student.student_id:restored},mock_mode=True)

    def test_handwritten_continuous_score_uses_public_half_up_rounding(self):
        from checker.ai import HandwrittenReviewer
        res=HandwrittenReviewer.normalize_review_result(dict(derivation_score=.745,
            clear_error_operators=['~'],error_evidence=['明确实质错误'],needs_review=False),
            self.student.student_id,self.student.name,1)
        self.assertEqual(res['derivation_base_score'],.75)
        self.assertEqual(res['suggested_score'],1.45)

    def test_explicit_legacy_csv_total_is_preserved_without_invented_problem_results(self):
        reporter=LabReporter('lab2',self.root/'explicit_legacy')
        path=reporter.output_dir/'lab2实验成绩单.csv'
        with path.open('w',encoding='utf-8-sig',newline='') as f:
            w=csv.DictWriter(f,fieldnames=['学号','总分','满分','建议扣分合计']);w.writeheader()
            w.writerow({'学号':self.student.student_id,'总分':'4.67','满分':'5','建议扣分合计':'.33'})
        g=reporter.load_existing_grades_from_files([self.student])[self.student.student_id]
        self.assertEqual(g.total_score,4.67);self.assertFalse(g.is_pending);self.assertTrue(g.needs_review)
        self.assertEqual(g.problem_results,{})
        files=reporter.generate_all_reports([self.student],{self.student.student_id:g})
        rows=list(csv.DictReader(files['detail_csv'].open(encoding='utf-8-sig')))
        self.assertEqual(len(rows),1)
        self.assertEqual(rows[0]['评分项'],'历史总分记录')
        self.assertNotIn('全部测试用例正确',files['roster_csv'].read_text(encoding='utf-8-sig'))

    def test_mock_reports_and_state_never_overwrite_formal_reports(self):
        j=ThirdLab();src=self.write('mode.json',json.dumps({'passed':[True]*10}))
        p=j.evaluate(self.student,[src],self.root/'build')
        g=StudentGrade(self.student,total_score=1,max_score=1,status='graded',problem_results={'99-1':p},comment=j.generate_comment(None))
        reporter=LabReporter('lab99',self.root/'modes')
        formal=reporter.generate_all_reports([self.student],{self.student.student_id:g})
        saved={k:v.read_bytes() for k,v in formal.items()}
        mocked=reporter.generate_all_reports([self.student],{self.student.student_id:g},mock_mode=True)
        self.assertTrue(all(v.name.startswith('mock_') for v in mocked.values()))
        for key,path in formal.items():self.assertEqual(path.read_bytes(),saved[key])
        state=reporter.load_existing_grades_from_files([self.student],mock_mode=False)
        self.assertEqual(state[self.student.student_id].total_score,1)

    def test_optional_palindrome_partial_versions_are_compared_without_early_stop(self):
        j=Lab2Judge()
        a=self.write('pal_more.c','#include <stdio.h>\nint main(){puts("5 7 11");}')
        b=self.write('pal_less.c','#include <stdio.h>\nint main(){puts("5");}')
        for sources in [[a,b],[b,a]]:
            p=j._grade_lab2_3(self.student,sources,self.root/'pal')
            self.assertEqual(p.score,0)
            self.assertEqual(len(p.candidate_evaluations),2)
            self.assertEqual(p.chosen_candidate['src_path'],str(a))
            self.assertEqual(p.chosen_candidate['passed_cases'],3)
            self.assertEqual(p.status,'partial')

    def test_single_pending_version_preserves_unrelated_confirmed_penalties(self):
        j=ThirdLab()
        src=self.write('single_pending.json',json.dumps({'passed':[True]*9+[False],'pending':True,'penalty':.1}))
        p=j.evaluate(self.student,[src],self.root/'build')
        self.assertTrue(p.is_pending)
        by_name={i.sub_item:i for i in p.check_items}
        self.assertIsNone(by_name['用例'].deduction)
        self.assertEqual(by_name['格式'].deduction,.1)
        self.assertFalse(by_name['格式'].needs_review)
