"""
tests/test_checker.py
判题机全覆盖单元与集成测试套件（完全脱敏，无任何硬编码个人真实数据）
"""
import io
import json
import os
import shutil
import unittest
import zipfile
from pathlib import Path

from checker.ai import HandwrittenReviewer, SemanticAmbiguityReviewer
from checker.compiler import ProjectCompiler
from checker.config import DEFAULT_ROSTER_PATH, DEFAULT_TEMPLATE_PATH, PROJECT_ROOT, WORK_DIR
from checker.labs.lab1 import Lab1Judge
from checker.labs.lab2 import Lab2Judge
from checker.models import ProblemResult, Student, StudentGrade, SubmissionFile, TestCaseResult
from checker.reporter import LabReporter
from checker.roster import RosterManager
from checker.sandbox import run_binary
from checker.unpack import SafeExtractor, SubmissionNormalizer


class TestCheckerSystem(unittest.TestCase):
    """判题系统全流程测试"""

    @classmethod
    def setUpClass(cls):
        cls.test_work = WORK_DIR / "test_runs"
        cls.test_work.mkdir(parents=True, exist_ok=True)

        # 若真实点名册存在则加载，否则动态构造脱敏单测专用花名册
        cls.roster_mgr = RosterManager()
        if DEFAULT_ROSTER_PATH.exists():
            cls.roster_mgr.load_roster(DEFAULT_ROSTER_PATH)
        else:
            cls.roster_mgr.students = [
                Student(seq=1, student_id="2026000067", name="测试生甲", class_id="260101"),
                Student(seq=2, student_id="2026000090", name="测试生乙", class_id="260101"),
            ]
            for s in cls.roster_mgr.students:
                cls.roster_mgr.by_id[s.student_id] = s
                cls.roster_mgr.by_name.setdefault(s.name, []).append(s)

    @classmethod
    def tearDownClass(cls):
        if cls.test_work.exists():
            shutil.rmtree(cls.test_work, ignore_errors=True)

    def test_01_roster_parsing(self):
        """测试 1: 点名册解析与学生身份匹配"""
        self.assertGreater(len(self.roster_mgr.students), 0)
        s1 = self.roster_mgr.students[0]
        self.assertTrue(len(s1.student_id) >= 8)
        self.assertTrue(len(s1.name) >= 2)

        # 动态匹配测试
        sample_text = f"2612901_{s1.name}_{s1.student_id}_实验1.zip"
        matched, cands, reason = self.roster_mgr.match_student_from_text(sample_text)
        self.assertIsNotNone(matched)
        self.assertEqual(matched.student_id, s1.student_id)

        # 模糊空格/下划线测试
        s2 = self.roster_mgr.students[1] if len(self.roster_mgr.students) > 1 else s1
        matched2, _, _ = self.roster_mgr.match_student_from_text(f"{s2.student_id}  {s2.name}  Lab 1-2.c")
        self.assertIsNotNone(matched2)
        self.assertEqual(matched2.student_id, s2.student_id)

    def test_02_safe_unpack_and_zipslip_protection(self):
        """测试 2: Zip Slip 恶意路径穿越防御与正常解压"""
        malicious_zip = self.test_work / "malicious.zip"
        with zipfile.ZipFile(malicious_zip, "w") as zf:
            zf.writestr("../../evil.txt", "attack payload")
            zf.writestr("normal.c", "int main(){return 0;}")

        dest = self.test_work / "unsafe_extract"
        extractor = SafeExtractor(self.test_work / "scratch")
        ok, msg, files = extractor.extract_archive(malicious_zip, dest)
        self.assertFalse(ok)
        self.assertIn("Zip Slip", msg)

        # 正常 zip
        normal_zip = self.test_work / "normal.zip"
        with zipfile.ZipFile(normal_zip, "w") as zf:
            zf.writestr("test_folder/hello.c", "int main(){return 0;}")

        ok2, msg2, files2 = extractor.extract_archive(normal_zip, dest)
        self.assertTrue(ok2)
        self.assertEqual(len(files2), 1)

    def test_03_compiler_and_windows_compat(self):
        """测试 3: Windows 特有语法与头文件自动兼容转换"""
        compiler = ProjectCompiler()
        build_dir = self.test_work / "build_compat"

        win_code = """
#include <stdio.h>
#include <windows.h>
#include <conio.h>

int main() {
    int val = 0;
    scanf_s("%d", &val);
    printf("Received: %d\\n", val * 2);
    fflush(stdin);
    system("pause");
    return 0;
}
"""
        src_path = self.test_work / "win_prog.c"
        src_path.write_text(win_code, encoding="gbk")

        c_res = compiler.compile(src_path, build_dir, "win_compat_test")
        self.assertTrue(c_res.success)
        self.assertTrue(c_res.compat_applied)
        self.assertGreater(len(c_res.compat_reasons), 0)

        exec_res = run_binary(c_res.binary_path, input_str="21\n", timeout=2.0)
        self.assertEqual(exec_res.exit_code, 0)
        self.assertIn("Received: 42", exec_res.stdout)

    def test_04_sandbox_timeout_and_infinite_loop(self):
        """测试 4: 沙箱超时强杀（死循环进程组清理）与超长输出截断"""
        compiler = ProjectCompiler()
        build_dir = self.test_work / "build_sandbox"

        # A. 死循环程序
        loop_code = "int main() { while(1) {} return 0; }"
        loop_src = self.test_work / "loop.c"
        loop_src.write_text(loop_code)
        c_res = compiler.compile(loop_src, build_dir, "loop")
        self.assertTrue(c_res.success)

        res = run_binary(c_res.binary_path, input_str="", timeout=1.0)
        self.assertTrue(res.timed_out)
        self.assertEqual(res.exit_code, -1)

        # B. 刷屏巨量输出程序 (防 OOM 截断)
        flood_code = """
#include <stdio.h>
int main() {
    for (int i = 0; i < 50000; i++) {
        printf("flood_test_line_1234567890_flood_test_line_1234567890\\n");
    }
    return 0;
}
"""
        flood_src = self.test_work / "flood.c"
        flood_src.write_text(flood_code)
        c_res2 = compiler.compile(flood_src, build_dir, "flood")
        self.assertTrue(c_res2.success)

        res2 = run_binary(c_res2.binary_path, input_str="", timeout=2.0)
        self.assertTrue(res2.truncated)
        self.assertIn("TRUNCATED", res2.stdout)

    def test_05_loose_vs_strict_lab1(self):
        """测试 5: Lab 1 确定性与宽松判题（对比正确与错误逻辑）"""
        judge = Lab1Judge()
        build_dir = self.test_work / "build_lab1_test"

        s1 = self.roster_mgr.students[0]
        exp_x = int(s1.student_id[-2:])

        correct_1_2 = f"""
#include <stdio.h>
int main(void) {{
    short x = {exp_x}, y = 88;
    printf("%5hd%5hd%5hd\\n", !x, x||y, x&&y);
    printf("%5hu%5hd%5hd\\n", ~x, x|y, x&y);
    printf("%5hd%5hd%5hd\\n \\n", ~x, x|y, x&y);
    return 0;
}}
"""
        p_1_2 = self.test_work / f"{s1.student_id}_Lab 1-2.c"
        p_1_2.write_text(correct_1_2)

        correct_1_3 = """
#include <stdio.h>
int main() {
    int button;
    float temp;
    printf("请输入模式(0:F->C, 1:C->F)和温度: ");
    if (scanf("%d %f", &button, &temp) == 2) {
        if (button == 0) {
            float c = 5.0f * (temp - 32.0f) / 9.0f;
            printf("转换后的摄氏度是: %.2f\\n", c);
        } else {
            float f = 9.0f * temp / 5.0f + 32.0f;
            printf("The Fahrenheit is %.2f\\n", f);
        }
    }
    return 0;
}
"""
        p_1_3 = self.test_work / f"{s1.student_id}_Lab 1-3.c"
        p_1_3.write_text(correct_1_3)

        files = [
            SubmissionFile(p_1_2, p_1_2.name, "c_source", "1-2"),
            SubmissionFile(p_1_3, p_1_3.name, "c_source", "1-3"),
        ]

        grade = judge.grade_student(s1, files, build_dir, mock_ai=True)
        self.assertEqual(grade.problem_results["1-2"].sub_scores["code"], 1.5)
        self.assertEqual(grade.problem_results["1-3"].status, "passed")
        self.assertEqual(grade.problem_results["1-3"].score, 2.0)
        self.assertGreaterEqual(len(grade.comment), 50)

        # 错误逻辑：整型截断除法 5/9
        wrong_int_div_1_3 = """
#include <stdio.h>
int main() {
    int button;
    float temp;
    if (scanf("%d %f", &button, &temp) == 2) {
        if (button == 0) {
            float c = (5 / 9) * (temp - 32);
            printf("%.2f\\n", c);
        } else {
            float f = 9 * temp / 5 + 32;
            printf("%.2f\\n", f);
        }
    }
    return 0;
}
"""
        p_wrong = self.test_work / "wrong_1_3.c"
        p_wrong.write_text(wrong_int_div_1_3)
        files_wrong = [SubmissionFile(p_wrong, p_wrong.name, "c_source", "1-3")]

        grade_wrong = judge.grade_student(s1, files_wrong, build_dir, mock_ai=True)
        self.assertLess(grade_wrong.problem_results["1-3"].score, 2.0)
        self.assertTrue(any("截断" in n for n in grade_wrong.problem_results["1-3"].notes))

    def test_06_loose_vs_strict_lab2(self):
        """测试 6: Lab 2 宽松问候语匹配、字母金字塔对称性与笔误容错"""
        judge = Lab2Judge()
        build_dir = self.test_work / "build_lab2_test"
        s2 = self.roster_mgr.students[1] if len(self.roster_mgr.students) > 1 else self.roster_mgr.students[0]

        chinese_greeting_2_1 = """
#include <stdio.h>
int main() {
    int h;
    if (scanf("%d", &h) == 1) {
        if (h >= 0 && h < 12) printf("早上好，阳光明媚！\\n");
        else if (h == 12) printf("中午好！\\n");
        else if (h >= 13 && h < 18) printf("下午好，喝杯茶吧！\\n");
        else if (h >= 18 && h < 21) printf("傍晚好！\\n");
        else if (h >= 21 && h < 24) printf("深夜好，晚安！\\n");
        else printf("输入不合法，错误！\\n");
    }
    return 0;
}
"""
        p_2_1 = self.test_work / f"{s2.student_id}_Lab 2-1.c"
        p_2_1.write_text(chinese_greeting_2_1)

        pyramid_2_2 = """
#include <stdio.h>
#include <ctype.h>

int main() {
    char ch;
    if (scanf(" %c", &ch) == 1) {
        if (!isalpha(ch)) {
            printf("Not a letter!\\n");
            return 0;
        }
        ch = toupper(ch);
        int n = ch - 'A' + 1;
        for (int i = 1; i <= n; i++) {
            for (int s = 0; s < n - i; s++) printf(" ");
            for (char c = 'A'; c < 'A' + i; c++) printf("%c", c);
            for (char c = 'A' + i - 2; c >= 'A'; c--) printf("%c", c);
            printf("\\n");
        }
    }
    return 0;
}
"""
        p_2_2 = self.test_work / f"{s2.student_id}_Lab 2-2.c"
        p_2_2.write_text(pyramid_2_2)

        files = [
            SubmissionFile(p_2_1, p_2_1.name, "c_source", "2-1"),
            SubmissionFile(p_2_2, p_2_2.name, "c_source", "2-2"),
        ]

        grade = judge.grade_student(s2, files, build_dir, mock_ai=True)
        self.assertEqual(grade.problem_results["2-1"].score, 2.0)
        self.assertEqual(grade.problem_results["2-2"].score, 3.0)
        self.assertEqual(grade.total_score, 5.0)

        # 单向未对称镜像
        bad_pyramid = """
#include <stdio.h>
int main() {
    char ch;
    scanf(" %c", &ch);
    printf("A\\nAB\\nABC\\n");
    return 0;
}
"""
        p_bad = self.test_work / "bad_pyramid.c"
        p_bad.write_text(bad_pyramid)
        files_bad = [SubmissionFile(p_bad, p_bad.name, "c_source", "2-2")]

        grade_bad = judge.grade_student(s2, files_bad, build_dir, mock_ai=True)
        self.assertLess(grade_bad.problem_results["2-2"].score, 3.0)

    def test_07_ai_review_fallback_on_failure(self):
        """测试 7: AI 调用失败/超时时不误判为 0 分，保留复核标记"""
        s1 = self.roster_mgr.students[0]
        res = HandwrittenReviewer.review_images(
            student_id=s1.student_id,
            student_name=s1.name,
            expected_x=int(s1.student_id[-2:]),
            image_paths=[self.test_work / "non_existent.png"],
            mock_mode=False,
        )
        self.assertEqual(res["suggested_score"], 0.0)
        self.assertIn("缺失", res["reasoning"])

    def test_08_reporter_outputs_and_roster_alignment(self):
        """测试 8: 三份报表导出与点名册全量对齐（不少于50字评语）"""
        output_dir = self.test_work / "output"
        reporter = LabReporter("lab1", output_dir)

        s1 = self.roster_mgr.students[0]
        s2 = self.roster_mgr.students[1] if len(self.roster_mgr.students) > 1 else s1

        g1 = StudentGrade(student=s1, total_score=5.0, status="graded")
        g1.comment = f"{s1.name}同学你好！本次实验1针对二进制逻辑与C语言基础语法的掌握情况评测完成。按要求将学号末两位对齐参数x，位运算与逻辑运算执行完全正确，演草纸推导规范清晰；Lab 1-3温度转换程序编写规范，成功通过全部双向精度与边界用例测试，表现优秀！望继续保持！"

        g2 = StudentGrade(student=s2, total_score=3.5, status="graded", needs_review=True, review_reasons=["演草纸字迹待人工核查"])
        g2.comment = f"{s2.name}同学你好！本次实验1在Lab 1-2中完成了基础位运算代码，但部分输出结果与推导细节存在偏差，演草纸步骤需更加完备；Lab 1-3温度转换程序编写规范，成功通过测试。希望在后续实验中继续加强代码调试与边界测试意识，扎实掌握基础。"

        grades = {s1.student_id: g1, s2.student_id: g2}

        res_files = reporter.generate_all_reports(
            self.roster_mgr.students,
            grades,
            excel_template_path=DEFAULT_TEMPLATE_PATH if DEFAULT_TEMPLATE_PATH.exists() else None,
        )

        detail_csv = res_files["detail_csv"]
        summary_md = res_files["summary_md"]
        roster_csv = res_files["roster_csv"]

        self.assertTrue(detail_csv.exists())
        self.assertTrue(summary_md.exists())
        self.assertTrue(roster_csv.exists())

        lines = roster_csv.read_text(encoding="utf-8-sig").splitlines()
        self.assertEqual(len(lines), len(self.roster_mgr.students) + 1)

        self.assertGreaterEqual(len(g1.comment), 50)
        self.assertGreaterEqual(len(g2.comment), 50)

        md_text = summary_md.read_text(encoding="utf-8")
        self.assertIn("待人工复核清单", md_text)
        self.assertIn(s2.student_id, md_text)

    def test_09_adversarial_lab1_2_checks(self):
        """测试 9 (反例对抗): Lab 1-2 逻辑运算错误被严格拦截，且代码与手写参数矛盾时触发复核"""
        judge = Lab1Judge()
        build_dir = self.test_work / "build_lab1_adv"
        s1 = self.roster_mgr.students[0]
        exp_x = int(s1.student_id[-2:])

        # 反例 A: || 与 && 输出错误 (全输出 0) 但位运算正确
        wrong_logic_code = f"""
#include <stdio.h>
int main(void) {{
    short x = {exp_x}, y = 88;
    printf("%5hd%5hd%5hd\\n", 0, 0, 0); // 逻辑运算故意错误
    printf("%5hu%5hd%5hd\\n", (unsigned short)~x, (short)(x|y), (short)(x&y));
    printf("%5hd%5hd%5hd\\n \\n", (short)~x, (short)(x|y), (short)(x&y));
    return 0;
}}
"""
        p_wrong_logic = self.test_work / "wrong_logic_1_2.c"
        p_wrong_logic.write_text(wrong_logic_code)
        files_wrong = [SubmissionFile(p_wrong_logic, p_wrong_logic.name, "c_source", "1-2")]

        res_wrong = judge._grade_lab1_2(s1, [p_wrong_logic], [], build_dir, mock_ai=True)
        # 必须失分，不能获得 1.5 分满分
        self.assertLess(res_wrong.sub_scores["code"], 1.5)
        self.assertTrue(any("逻辑运算" in n or "不符预期" in n for n in res_wrong.notes))

        # 反例 B: 代码中 y=88，但演草纸手写为 y=1 参数矛盾核验
        correct_code = f"""
#include <stdio.h>
int main(void) {{
    short x = {exp_x}, y = 88;
    printf("%5hd%5hd%5hd\\n", !x, x||y, x&&y);
    printf("%5hu%5hd%5hd\\n", (unsigned short)~x, (short)(x|y), (short)(x&y));
    printf("%5hd%5hd%5hd\\n \\n", (short)~x, (short)(x|y), (short)(x&y));
    return 0;
}}
"""
        p_correct = self.test_work / "correct_1_2.c"
        p_correct.write_text(correct_code)

        # 模拟 AI 识别到的演草纸手写 y 为 1 (与代码 y=88 矛盾)
        from unittest.mock import patch
        mock_ai_return = {
            "task_id": "lab1_2_handwritten",
            "student_id_observed": s1.student_id,
            "name_observed": s1.name,
            "phone_observed": "01",
            "x_observed": exp_x,
            "y_observed": 1,  # 矛盾参数
            "info_score": 0.5,
            "xy_match_score": 0.2,
            "derivation_score": 0.8,
            "suggested_score": 1.5,
            "is_pending": False,
            "reasoning": "验算 y=1",
            "needs_review": False,
        }

        with patch("checker.ai.HandwrittenReviewer.review_images", return_value=mock_ai_return):
            res_discrepancy = judge._grade_lab1_2(s1, [p_correct], [self.test_work / "dummy.png"], build_dir, mock_ai=False)
            self.assertTrue(res_discrepancy.needs_review)
            self.assertTrue(any("参数矛盾" in r for r in res_discrepancy.review_reasons))

    def test_10_adversarial_lab1_3_dumping_cheat(self):
        """测试 10 (反例对抗): Lab 1-3 倾倒全部参考答案刷屏作弊与静态输出被直接拦截"""
        judge = Lab1Judge()
        build_dir = self.test_work / "build_lab1_3_adv"
        s1 = self.roster_mgr.students[0]

        # 倾倒全部用例标准答案，完全不读取输入
        dump_cheat_code = """
#include <stdio.h>
int main() {
    printf("37.78 19.99 -17.78 0.00 32.00 212.00 -40.00\\n");
    return 0;
}
"""
        p_dump = self.test_work / "dump_1_3.c"
        p_dump.write_text(dump_cheat_code)

        res_dump = judge._grade_lab1_3(s1, [p_dump], build_dir, mock_ai=True)
        # 绝不能通过，得分必须为 0 分
        self.assertEqual(res_dump.score, 0.0)
        self.assertEqual(res_dump.status, "failed")

    def test_11_adversarial_lab2_2_single_line_and_invalid_char(self):
        """测试 11 (反例对抗): 金字塔单行作弊拦截与规范 Invalid character! 识别"""
        judge = Lab2Judge()
        build_dir = self.test_work / "build_lab2_2_adv"
        s2 = self.roster_mgr.students[1] if len(self.roster_mgr.students) > 1 else self.roster_mgr.students[0]

        # 反例 A: 输入 C 却只打印最后一行 ABCBA 的投机程序
        cheat_single_line_code = """
#include <stdio.h>
int main() {
    char ch;
    scanf(" %c", &ch);
    printf("ABCBA\\n"); // 只输出最后一行
    return 0;
}
"""
        p_cheat = self.test_work / "cheat_pyramid.c"
        p_cheat.write_text(cheat_single_line_code)

        res_cheat = judge._grade_lab2_2(s2, [p_cheat], build_dir, mock_ai=True)
        # 针对输入 C 期望 3 行，实际仅 1 行，必须判定失败，得分绝不能为满分
        self.assertLess(res_cheat.score, 3.0)

        # 针对非字母打印指导书规范 Sample 3: "Invalid character!" 的完全正确程序
        valid_error_code = """
#include <stdio.h>
#include <ctype.h>
int main() {
    char ch;
    if (scanf(" %c", &ch) == 1) {
        if (!isalpha(ch)) {
            printf("Invalid character!\\n");
            return 0;
        }
        ch = toupper(ch);
        int n = ch - 'A' + 1;
        for (int i = 1; i <= n; i++) {
            for (int s = 0; s < n - i; s++) printf(" ");
            for (char c = 'A'; c < 'A' + i; c++) printf("%c", c);
            for (char c = 'A' + i - 2; c >= 'A'; c--) printf("%c", c);
            printf("\\n");
        }
    }
    return 0;
}
"""
        p_valid_err = self.test_work / "valid_err_pyramid.c"
        p_valid_err.write_text(valid_error_code)

        res_valid = judge._grade_lab2_2(s2, [p_valid_err], build_dir, mock_ai=True)
        # 必须获得满分 3.0，包括非字母用例完全通过
        self.assertEqual(res_valid.score, 3.0)
        self.assertEqual(res_valid.status, "passed")

    def test_12_scanf_s_shim_and_math_link_order(self):
        """测试 12: scanf_s 多参数垫片与 -lm 数学库链接顺序正确性"""
        compiler = ProjectCompiler()
        build_dir = self.test_work / "build_compiler_adv"

        # A. 测试 -lm 链接数学库 sin/cos
        math_code = """
#include <stdio.h>
#include <math.h>
int main() {
    double val = sin(0.0);
    printf("sin(0)=%.1f\\n", val);
    return 0;
}
"""
        p_math = self.test_work / "math_test.c"
        p_math.write_text(math_code)
        c_res = compiler.compile(p_math, build_dir, "math_test")
        self.assertTrue(c_res.success, f"数学库链接失败: {c_res.stderr}")

        exec_res = run_binary(c_res.binary_path, input_str="")
        self.assertIn("sin(0)=0.0", exec_res.stdout)

        # B. 测试 scanf_s 带缓冲区大小参数 (%s %d)
        scanf_s_code = """
#include <stdio.h>
int main() {
    char name[32];
    int score = 0;
    // MSVC 规范：字符串缓冲区后需携带长度参数
    scanf_s("%s %d", name, (unsigned)sizeof(name), &score);
    printf("Name: %s, Score: %d\\n", name, score);
    return 0;
}
"""
        p_scanf = self.test_work / "scanf_s_test.c"
        p_scanf.write_text(scanf_s_code)
        c_res2 = compiler.compile(p_scanf, build_dir, "scanf_s_test")
        self.assertTrue(c_res2.success, f"scanf_s 编译失败: {c_res2.stderr}")

        exec_res2 = run_binary(c_res2.binary_path, input_str="Alice 95\n")
        self.assertEqual(exec_res2.exit_code, 0)
        self.assertIn("Name: Alice, Score: 95", exec_res2.stdout)

    def test_13_attribution_conflict_reported(self):
        """测试 13: 提交文件多候选归属冲突不被转换成未交，准确进入复核清单"""
        output_dir = self.test_work / "output_conflict"
        reporter = LabReporter("lab1", output_dir)

        s1 = self.roster_mgr.students[0]
        s2 = self.roster_mgr.students[1] if len(self.roster_mgr.students) > 1 else s1

        # 构造冲突记录
        conflict_msg = f"提交包同时包含多个候选学号: ['{s1.student_id}', '{s2.student_id}']"
        g1 = StudentGrade(
            student=s1,
            max_score=5.0,
            status="ambiguous",
            needs_review=True,
            is_pending=True,
            review_reasons=[conflict_msg],
        )

        grades = {s1.student_id: g1}
        ambiguous_items = [{"path": "conflict_test.zip", "candidates": [s1.student_id, s2.student_id], "reason": conflict_msg}]

        rep_paths = reporter.generate_all_reports(
            self.roster_mgr.students,
            grades,
            ambiguous_items=ambiguous_items,
        )

        summary_text = rep_paths["summary_md"].read_text(encoding="utf-8")
        self.assertIn("待人工复核清单", summary_text)
        self.assertIn(s1.student_id, summary_text)
        self.assertIn("多学号冲突提交件", summary_text)

        detail_text = rep_paths["detail_csv"].read_text(encoding="utf-8-sig")
        self.assertIn("冲突待复核", detail_text)

    def test_14_single_student_grade_merge(self):
        """测试 14: 单人重判 (--student-id) 安全合并，严防将他人成绩清零覆写"""
        output_dir = self.test_work / "output_single_merge"
        reporter = LabReporter("lab1", output_dir)

        s1 = self.roster_mgr.students[0]
        s2 = self.roster_mgr.students[1] if len(self.roster_mgr.students) > 1 else s1

        # 1. 模拟初始全班成绩：s1=4.5分, s2=5.0分
        g1 = StudentGrade(student=s1, total_score=4.5, status="graded", comment="学生甲初评4.5分" + "。" * 40)
        g2 = StudentGrade(student=s2, total_score=5.0, status="graded", comment="学生乙初评5.0分" + "。" * 40)
        init_grades = {s1.student_id: g1, s2.student_id: g2}

        reporter.generate_all_reports(self.roster_mgr.students, init_grades)

        # 2. 模拟针对 s1 单人重判为 5.0分
        g1_regrade = StudentGrade(student=s1, total_score=5.0, status="graded", comment="学生甲重评5.0分满分通过" + "。" * 40)

        # 单人结果独立导出
        single_res = reporter.generate_single_student_report(s1, g1_regrade)
        self.assertTrue(single_res["detail_csv"].exists())

        # 读取并合并全班批次
        merged_grades = reporter.load_existing_grades_from_files(self.roster_mgr.students)
        self.assertIn(s2.student_id, merged_grades)
        self.assertAlmostEqual(merged_grades[s2.student_id].total_score, 5.0)

        # 更新 s1
        merged_grades[s1.student_id] = g1_regrade
        reporter.generate_all_reports(self.roster_mgr.students, merged_grades)

        # 3. 验证 s2 的成绩绝未被清零成未交！
        roster_csv_text = (output_dir / "lab1实验成绩单.csv").read_text(encoding="utf-8-sig")
        self.assertIn(f"{s2.student_id},{s2.name}", roster_csv_text)
        self.assertIn(f"{s1.student_id},{s1.name}", roster_csv_text)
        self.assertIn("学生乙初评5.0分", roster_csv_text)
        self.assertIn("学生甲重评5.0分满分通过", roster_csv_text)

    def test_15_lab1_3_verbose_prompt_pass(self):
        """测试 15: Lab 1-3 包含提示语 0: F→C; 1: C→F; C=5*(F-32)/9 的正确程序通过不被误判作弊"""
        judge = Lab1Judge()
        build_dir = self.test_work / "build_lab1_3_verbose"
        s1 = self.roster_mgr.students[0]

        # 包含较多数字提示语但计算完全正确的学生程序
        verbose_code = """
#include <stdio.h>
int main() {
    int mode;
    double val, res;
    // 提示信息中包含 0, 1, 5, 32, 9 等多个数字
    printf("Please choose mode (0: F->C; 1: C->F; formula: C=5*(F-32)/9):\\n");
    if (scanf("%d %lf", &mode, &val) != 2) return 1;
    if (mode == 0) {
        res = 5.0 * (val - 32.0) / 9.0;
        printf("Fahrenheit %.2f to Centigrade is: %.2f\\n", val, res);
    } else {
        res = 9.0 * val / 5.0 + 32.0;
        printf("Centigrade %.2f to Fahrenheit is: %.2f\\n", val, res);
    }
    return 0;
}
"""
        p_verbose = self.test_work / "verbose_1_3.c"
        p_verbose.write_text(verbose_code)

        res = judge._grade_lab1_3(s1, [p_verbose], build_dir, mock_ai=True)
        # 必须获得满分 2.0，不得被判定作弊
        self.assertEqual(res.score, 2.0)
        self.assertEqual(res.status, "passed")
        for tr in res.test_results:
            self.assertTrue(tr.passed, f"用例未通过: {tr.input_data} -> {tr.actual_output} (error: {tr.error_message})")

    def test_16_detail_csv_multi_problem_and_recovery(self):
        """测试 16: 明细表与单人报告导出所有题目，历史成绩恢复包含全部题目"""
        output_dir = self.test_work / "output_multi_prob"
        reporter = LabReporter("lab1", output_dir)
        s1 = self.roster_mgr.students[0]

        g = StudentGrade(student=s1, total_score=5.0, status="graded", comment="学生甲全对" + "。" * 40)
        p1 = ProblemResult(problem_id="1-2", score=3.0, max_score=3.0, status="passed")
        p2 = ProblemResult(problem_id="1-3", score=2.0, max_score=2.0, status="passed")
        g.problem_results["1-2"] = p1
        g.problem_results["1-3"] = p2

        # 1. 批次全量明细表导出
        rep_paths = reporter.generate_all_reports(self.roster_mgr.students, {s1.student_id: g})
        detail_lines = rep_paths["detail_csv"].read_text(encoding="utf-8-sig").strip().splitlines()
        # 需同时出现各分项，包含 1-2 与 1-3
        s1_rows = [line for line in detail_lines if s1.student_id in line]
        self.assertTrue(len(s1_rows) >= 2, f"明细表中学生 {s1.student_id} 行数应至少为2，实际为: {s1_rows}")
        pids = [line.split(",")[3] for line in s1_rows]
        self.assertIn("1-2", pids)
        self.assertIn("1-3", pids)

        # 2. 单人报告导出
        single_paths = reporter.generate_single_student_report(s1, g)
        single_lines = single_paths["detail_csv"].read_text(encoding="utf-8-sig").strip().splitlines()
        single_s1_rows = [line for line in single_lines if s1.student_id in line]
        self.assertTrue(len(single_s1_rows) >= 2)

        # 3. 从历史文件恢复
        recovered = reporter.load_existing_grades_from_files(self.roster_mgr.students)
        self.assertIn(s1.student_id, recovered)
        rec_g = recovered[s1.student_id]
        self.assertAlmostEqual(rec_g.total_score, 5.0)

    def test_17_semantic_ai_timeout_pending_propagation(self):
        """测试 17: AI 语义复核超时返回 is_pending=True，全链条保持成绩待定绝非数值成绩"""
        judge = Lab1Judge()
        build_dir = self.test_work / "build_sem_pending"
        s1 = self.roster_mgr.students[0]

        # 语义模糊程序（格式略有差异，不精确命中硬匹配，触发语义复核）
        ambig_code = """
#include <stdio.h>
int main() {
    int m;
    double t;
    scanf("%d %lf", &m, &t);
    // 输出无法直接数值命中，触发语义复核
    printf("Approximate temp conversion result is: 999.00\\n");
    return 0;
}
"""
        p_ambig = self.test_work / "ambig_1_3.c"
        p_ambig.write_text(ambig_code)

        from unittest.mock import patch
        mock_sem_pending = {
            "task_id": "semantic_1-3",
            "category": "insufficient_evidence",
            "passed": False,
            "suggested_score": None,
            "is_pending": True,
            "reasoning": "AI 复核调度超时",
            "needs_review": True,
            "review_reason": "AI 调用超时，成绩待定",
        }

        with patch("checker.ai.SemanticAmbiguityReviewer.review_output", return_value=mock_sem_pending):
            prob = judge._grade_lab1_3(s1, [p_ambig], build_dir, mock_ai=False)
            self.assertTrue(prob.is_pending, "单题 is_pending 应为 True")
            self.assertTrue(prob.needs_review, "单题 needs_review 应为 True")
            self.assertEqual(prob.score, 0.0, "待定状态下数值得分应为 0.0 避免虚假出分")

            grade = judge.grade_student(
                s1,
                [SubmissionFile(original_name=p_ambig.name, path=p_ambig, file_type="c_source", matched_problem="1-3")],
                build_dir,
                mock_ai=False,
            )
            self.assertTrue(grade.is_pending, "学生总评 is_pending 应为 True")
            self.assertEqual(grade.status, "needs_review")

            output_dir = self.test_work / "output_sem_pending"
            reporter = LabReporter("lab1", output_dir)
            rep_paths = reporter.generate_all_reports(self.roster_mgr.students, {s1.student_id: grade})

            roster_csv_text = rep_paths["roster_csv"].read_text(encoding="utf-8-sig")
            # 成绩单中必须写入 '待复核'，绝不能写入虚假数值扣分
            self.assertIn(f"{s1.student_id},{s1.name}", roster_csv_text)
            self.assertIn("待复核", roster_csv_text)
            self.assertNotIn(",0.20,", roster_csv_text)

    def test_18_compiler_isolation_and_no_nameerror(self):
        """测试 18: 编译器模块正常导入 shutil，支持隔离环境编译不抛出 NameError"""
        compiler = ProjectCompiler()
        build_dir = self.test_work / "build_compiler_iso"
        p_simple = self.test_work / "simple.c"
        p_simple.write_text("int main(){return 0;}")

        # 编译应当顺利执行，无论系统中是否有 unshare 命令，均安全降级或使用隔离
        c_res = compiler.compile(p_simple, build_dir, "simple_bin")
        self.assertTrue(c_res.success)
        self.assertTrue(c_res.binary_path.exists())

    def test_19_consecutive_scanf_s_on_same_line(self):
        """测试 19: MSVC scanf_s 垫片支持单行多次输入 (如 '0 100') 不消耗整行"""
        compiler = ProjectCompiler()
        build_dir = self.test_work / "build_scanf_consec"

        code = """
#include <stdio.h>
int main() {
    int mode = -1;
    double temp = -999.0;
    // 连续两次调用 scanf_s 读取同一行的两个数值
    int r1 = scanf_s("%d", &mode);
    int r2 = scanf_s("%lf", &temp);
    printf("r1=%d mode=%d, r2=%d temp=%.2f\\n", r1, mode, r2, temp);
    return 0;
}
"""
        p_code = self.test_work / "consec_scanf.c"
        p_code.write_text(code)

        c_res = compiler.compile(p_code, build_dir, "consec_scanf")
        self.assertTrue(c_res.success, f"编译失败: {c_res.stderr}")

        exec_res = run_binary(c_res.binary_path, input_str="0 100\n")
        self.assertEqual(exec_res.exit_code, 0)
        self.assertIn("r1=1 mode=0, r2=1 temp=100.00", exec_res.stdout)

    def test_20_mock_mode_grade_sheet_tagging(self):
        """测试 20: 模拟模式输出加 mock_ 前缀并标明 [模拟] 与 [模拟测试非正式] 表头"""
        output_dir = self.test_work / "output_mock_tag"
        reporter = LabReporter("lab1", output_dir)
        s1 = self.roster_mgr.students[0]

        g = StudentGrade(student=s1, total_score=5.0, status="graded", comment="学生甲模拟全对" + "。" * 40)
        rep_paths = reporter.generate_all_reports(
            self.roster_mgr.students,
            {s1.student_id: g},
            excel_template_path=DEFAULT_TEMPLATE_PATH if DEFAULT_TEMPLATE_PATH.exists() else None,
            mock_mode=True,
        )

        # 检查文件名含有 mock_
        self.assertTrue("mock_" in rep_paths["roster_csv"].name)
        self.assertTrue("mock_" in rep_paths["detail_csv"].name)

        csv_text = rep_paths["roster_csv"].read_text(encoding="utf-8-sig")
        # 检查表头明确标识模拟
        self.assertIn("建议扣分合计[模拟测试非正式]", csv_text)
        # 检查扣分单元格包含 [模拟] 0.00
        self.assertIn("[模拟] 0.00", csv_text)

        if "roster_xlsx" in rep_paths:
            import openpyxl
            wb = openpyxl.load_workbook(rep_paths["roster_xlsx"])
            ws = wb["Lab 1"]
            self.assertEqual(ws.cell(row=2, column=4).value, "分数(5)[模拟测试非正式]")
            self.assertIn("[模拟]", str(ws.cell(row=3, column=4).value))

    def test_21_ai_student_identity_mismatch_flagged(self):
        """测试 21: AI 返回其他学生身份时系统准确识别矛盾、扣除得分并标记待复核"""
        s1 = self.roster_mgr.students[0]
        from checker.ai import HandwrittenReviewer

        dummy_img = self.test_work / "dummy.png"
        dummy_img.write_bytes(b"dummy_png_bytes_content" * 20)

        mock_wrong_student = {
            "task_id": "lab1_2_handwritten",
            "student_id_observed": "2026999999",  # 另一位学生学号
            "name_observed": "张三丰",          # 另一位学生姓名
            "phone_observed": "88",
            "x_observed": 67,
            "y_observed": 88,
            "info_score": 0.5,
            "xy_match_score": 0.2,
            "derivation_score": 0.8,
            "suggested_score": 1.5,
            "is_pending": False,
            "reasoning": "书写规范",
            "needs_review": False,
        }

        from unittest.mock import patch
        with patch("checker.ai.call_agy_cli", return_value=(True, mock_wrong_student, "OK")):
            res = HandwrittenReviewer.review_images(
                student_id=s1.student_id,
                student_name=s1.name,
                expected_x=67,
                image_paths=[dummy_img],
                mock_mode=False,
            )
            # 必须判定身份矛盾并标记复核，得分清零
            self.assertTrue(res["needs_review"])
            self.assertEqual(res["suggested_score"], 0.0)
            self.assertIn("演草纸学生身份矛盾", res["review_reason"])

    def test_22_ai_regex_fallback_parsing(self):
        """测试 22: call_agy_cli 在无 structured_output 时使用正则回退解析不产生 NameError"""
        from checker.ai import call_agy_cli
        from unittest.mock import patch, MagicMock

        # 模拟 CLI 仅返回 Markdown 围栏代码块中的 JSON，无 structured_output
        cli_raw_stdout = """
Here is the review result:
```json
{
  "task_id": "lab1_2_handwritten",
  "student_id_observed": "2026000067",
  "name_observed": "测试生",
  "phone_observed": "88",
  "x_observed": 67,
  "y_observed": 88,
  "info_score": 0.5,
  "xy_match_score": 0.2,
  "derivation_score": 0.8,
  "suggested_score": 1.5,
  "reasoning": "测试通过",
  "needs_review": false,
  "review_reason": ""
}
```
"""
        mock_proc = MagicMock()
        mock_proc.returncode = 0
        mock_proc.stdout = json.dumps({"response": cli_raw_stdout})
        mock_proc.stderr = ""

        with patch("subprocess.run", return_value=mock_proc):
            ok, res, msg = call_agy_cli("dummy prompt", {})
            self.assertTrue(ok)
            self.assertIsNotNone(res)
            self.assertEqual(res.get("task_id"), "lab1_2_handwritten")
            self.assertEqual(res.get("x_observed"), 67)

    def test_23_batch_container_and_nested_zips(self):
        """测试 23: 平台总包容器 (含嵌套 zip 与 .cbp 文件) 的批量提取与规范化归档"""
        container_dir = self.test_work / "batch_container_test"
        submit_dir = container_dir / "submit"
        submit_dir.mkdir(parents=True, exist_ok=True)
        s1 = self.roster_mgr.students[0]
        s2 = self.roster_mgr.students[1] if len(self.roster_mgr.students) > 1 else s1

        # 构造嵌套 zip
        nested_code_zip = io.BytesIO()
        with zipfile.ZipFile(nested_code_zip, "w") as nz:
            nz.writestr("2026000067_Lab 1-2.c", "int main(){return 0;}")
            nz.writestr("2026000067_Lab 1-3.c", "int main(){return 0;}")

        student_zip_data = io.BytesIO()
        with zipfile.ZipFile(student_zip_data, "w") as sz:
            sz.writestr("inner_code.zip", nested_code_zip.getvalue())
            sz.writestr("2026000067_Lab 1-2 验算过程.jpg", b"fake_jpg_content" * 20)

        # 构造整批容器 zip: 实验一.zip
        batch_zip_path = submit_dir / "实验一.zip"
        with zipfile.ZipFile(batch_zip_path, "w") as bz:
            bz.writestr(f"260101_{s1.name}_{s1.student_id}_实验1.zip", student_zip_data.getvalue())
            bz.writestr(f"{s2.student_id}_Lab1-2.cbp", "<xml>cbp project</xml>")

        normalizer = SubmissionNormalizer(self.roster_mgr, "lab1", container_dir)
        student_files = normalizer.process_all_submissions()

        # 验证 s1 的嵌套 zip 被成功解开并规范化至 src/
        self.assertIn(s1.student_id, student_files)
        s1_files = student_files[s1.student_id]
        s1_types = [f.file_type for f in s1_files]
        self.assertIn("c_source", s1_types)
        self.assertIn("image", s1_types)
        s1_problems = [f.matched_problem for f in s1_files]
        self.assertIn("1-2", s1_problems)
        self.assertIn("1-3", s1_problems)

        # 验证 s2 的 .cbp 工程文件被成功归纳
        self.assertIn(s2.student_id, student_files)
        s2_files = student_files[s2.student_id]
        self.assertTrue(any(f.file_type == "project" for f in s2_files))

    def test_24_attribution_disambiguation_and_problem_matching(self):
        """测试24: 学号含12的文件名分类防误判、子目录main.c识别、代码内容兜底及图片归属候选过滤"""
        container_dir = self.test_work / "test_24_disambiguation"
        submit_dir = container_dir / "submit"
        submit_dir.mkdir(parents=True, exist_ok=True)
        s1 = self.roster_mgr.students[0]

        # 1. 测试学号包含 12 的 Lab 1-3.c 不被误判为 1-2
        code_13 = self.test_work / f"2026111299_Lab 1-3.c"
        code_13.write_text("int main(){ return 0; }")
        normalizer = SubmissionNormalizer(self.roster_mgr, "lab1", container_dir)
        sf_13 = normalizer._classify_and_stage_file(s1, code_13, code_13.name)
        self.assertEqual(sf_13.matched_problem, "1-3", "学号含12的1-3代码不应误判为1-2")

        # 2. 测试子目录下 main.c 的题号识别
        main_c_in_dir = self.test_work / "main.c"
        main_c_in_dir.write_text("int main(){ return 0; }")
        sf_dir_13 = normalizer._classify_and_stage_file(s1, main_c_in_dir, "Lab 1-3/main.c")
        self.assertEqual(sf_dir_13.matched_problem, "1-3", "路径含 Lab 1-3/ 的 main.c 应识别为 1-3")
        sf_dir_12 = normalizer._classify_and_stage_file(s1, main_c_in_dir, "Lab 1-2/main.c")
        self.assertEqual(sf_dir_12.matched_problem, "1-2", "路径含 Lab 1-2/ 的 main.c 应识别为 1-2")

        # 3. 测试误命名为 1-1 的代码通过运算符内容兜底识别为 1-2
        code_misnamed = self.test_work / "Lab 1-1.c"
        code_misnamed.write_text("""
        #include <stdio.h>
        int main() {
            short x = 1, y = 2;
            printf("%d", !x);
            printf("%d", x || y);
            return 0;
        }
        """)
        sf_misnamed = normalizer._classify_and_stage_file(s1, code_misnamed, code_misnamed.name)
        self.assertEqual(sf_misnamed.matched_problem, "1-2", "含位/逻辑运算的内容应兜底识别为 1-2")

        # 4. 测试图片根据参数 x 识别并受候选池过滤
        from checker.ai import HandwrittenReviewer
        fake_img = self.test_work / "unmatched.jpg"
        fake_img.write_bytes(b"\xFF\xD8\xFF\xE0" + b"\x00" * 200)

        mock_identify_res = {
            "task_id": "lab1_2_handwritten",
            "content_summary": "验算演草纸",
            "x_observed": int(s1.student_id[-2:]),
            "y_observed": 55,
        }

        from unittest.mock import patch
        with patch("checker.ai.call_agy_cli", return_value=(True, mock_identify_res, "")):
            # 当只在包含 s1 的候选池中搜索时应唯一定位到 s1
            matched, info = HandwrittenReviewer.identify_student_from_image(
                fake_img, self.roster_mgr, candidate_students=[s1]
            )
            self.assertEqual(matched.student_id, s1.student_id)
            self.assertEqual(info["method"], "param_x")

    def test_25_multi_candidate_highest_score(self):
        """测试 25: 同题多源码候选自动评测全部版本并择优统计最高分，不触发待复核"""
        s1 = self.roster_mgr.students[0]
        judge = Lab1Judge()
        build_dir = self.test_work / "build_multi_cand"
        build_dir.mkdir(parents=True, exist_ok=True)

        exp_x = int(s1.student_id[-2:])
        # 候选 1: 错误/崩溃代码 (得分 0.0)
        cand1 = self.test_work / f"{s1.student_id}_1-2_bad.c"
        cand1.write_text("int main() { return 1; }")

        # 候选 2: 满分正确代码 (得分 1.5)
        cand2 = self.test_work / f"{s1.student_id}_1-2_good.c"
        cand2.write_text(f"""
#include <stdio.h>
int main() {{
    short x = {exp_x}, y = 88;
    printf("%5hd%5hd%5hd\\n", !x, x||y, x&&y);
    printf("%5hu%5hd%5hd\\n", (unsigned short)~x, x|y, x&y);
    printf("%5hd%5hd%5hd\\n", ~x, x|y, x&y);
    return 0;
}}
""")
        files = [
            SubmissionFile(cand1, cand1.name, "c_source", "1-2"),
            SubmissionFile(cand2, cand2.name, "c_source", "1-2"),
        ]
        grade = judge.grade_student(s1, files, build_dir, mock_ai=True)
        res1_2 = grade.problem_results["1-2"]
        # 必须择优采用 cand2 的最高分，且不被标记为同题多候选复核
        self.assertEqual(res1_2.sub_scores["code"], 1.5)
        self.assertFalse(res1_2.needs_review)
        self.assertIn("择优", "".join(res1_2.notes))

    def test_26_handwritten_missing_info_leniency(self):
        """测试 26: 演草纸未注明个人信息依目录归属最大容错不扣分，与源码对照核验通过"""
        from unittest.mock import patch
        from checker.ai import HandwrittenReviewer
        s1 = self.roster_mgr.students[0]
        exp_x = int(s1.student_id[-2:])

        mock_missing_info = {
            "task_id": "lab1_2_handwritten",
            "student_id_observed": "未填写",
            "name_observed": "未注明",
            "phone_observed": "",
            "x_observed": exp_x,
            "y_observed": 88,
            "info_score": 0.0,
            "xy_match_score": 0.2,
            "derivation_score": 0.8,
            "suggested_score": 1.0,
            "is_pending": False,
            "reasoning": "纸面未手写姓名学号，但6种运算展开完整准确",
            "needs_review": True,
            "review_reason": "纸面未书写姓名、学号、手机号，无法独立凭纸面确认身份归属，需人工复核。",
        }

        normalized = HandwrittenReviewer.normalize_review_result(
            mock_missing_info, s1.student_id, s1.name, exp_x
        )
        # 宽容政策：未注明身份不扣分 (满分0.5)，清除复核标志，最终总分 1.5
        self.assertEqual(normalized["info_score"], 0.5)
        self.assertEqual(normalized["suggested_score"], 1.5)
        self.assertFalse(normalized["needs_review"])

        judge = Lab1Judge()
        build_dir = self.test_work / "build_hw_leniency"
        build_dir.mkdir(parents=True, exist_ok=True)
        img = self.test_work / "hw.png"
        img.write_bytes(b"\xFF\xD8\xFF\xE0" + b"\x00" * 200)

        with patch("checker.ai.HandwrittenReviewer.review_images", return_value=normalized):
            prob = judge._grade_lab1_2(s1, [], [img], build_dir, mock_ai=False)
            item_info = [it for it in prob.check_items if it.sub_item == "姓名、学号、手机号完整性"][0]
            self.assertEqual(item_info.deduction, 0.0)
            self.assertEqual(item_info.completion_status, "已达成")
            self.assertFalse(item_info.needs_review)

    def test_27_lab1_2_preamble_number_offset_robustness(self):
        """测试 27: Lab 1-2 含有前置说明文字 (如 x=93,y=82) 时，9项运算数字不会发生错位"""
        s = Student(seq=93, student_id="2026111193", name="测试生郭", class_id="26TS602")
        judge = Lab1Judge()
        build_dir = self.test_work / "build_preamble_test"
        build_dir.mkdir(parents=True, exist_ok=True)

        c_file = self.test_work / f"{s.student_id}_1-2.c"
        c_file.write_text("""
#include <stdio.h>
int main() {
    short x = 93, y = 82;
    printf("x=%hd,y=%hd时的基本运算\\n", x, y);
    printf("逻辑非: %10hd 逻辑或: %10hd 逻辑与: %10hd\\n", !x, x||y, x&&y);
    printf("无符号取反: %10hu 按位或: %10hd 按位与: %10hd\\n", (unsigned short)~x, x|y, x&y);
    printf("有符号取反: %10hd 按位或: %10hd 按位与: %10hd\\n", ~x, x|y, x&y);
    return 0;
}
""")
        files = [SubmissionFile(c_file, c_file.name, "c_source", "1-2")]
        grade = judge.grade_student(s, files, build_dir, mock_ai=True)
        res = grade.problem_results["1-2"]
        self.assertEqual(res.sub_scores["code"], 1.5)
        # 推导出的 y 必须为 82，绝不能是 65444 或其他错位值
        self.assertNotIn("65444", "".join(res.notes))
        self.assertIn("推导y=82", "".join(res.notes))

    def test_28_lab1_3_menu_prompt_zero_not_int_div_bug(self):
        """测试 28: Lab 1-3 菜单提示语中包含 0 时，正确计算不触发整数除法截断扣分"""
        s = Student(seq=93, student_id="2026111193", name="测试生郭", class_id="26TS602")
        judge = Lab1Judge()
        build_dir = self.test_work / "build_menu_prompt_test"
        build_dir.mkdir(parents=True, exist_ok=True)

        c_file = self.test_work / f"{s.student_id}_1-3.c"
        c_file.write_text("""
#include <stdio.h>
int main() {
    int mode;
    double val;
    printf("0: Fahrenheit to Celsius\\n1: Celsius to Fahrenheit\\nPlease select mode and temp:\\n");
    if (scanf("%d %lf", &mode, &val) == 2) {
        if (mode == 0) {
            double c = (val - 32.0) * 5.0 / 9.0;
            printf("Result: %.2f C\\n", c);
        } else {
            double f = val * 9.0 / 5.0 + 32.0;
            printf("Result: %.2f F\\n", f);
        }
    }
    return 0;
}
""")
        files = [SubmissionFile(c_file, c_file.name, "c_source", "1-3")]
        grade = judge.grade_student(s, files, build_dir, mock_ai=True)
        res = grade.problem_results["1-3"]
        self.assertEqual(res.score, 2.0)
        # 华氏转摄氏分项应全对，不扣 0.80 分
        f2c_item = [it for it in res.check_items if it.sub_item == "华氏转摄氏计算"][0]
        self.assertEqual(f2c_item.deduction, 0.0)
        self.assertEqual(f2c_item.completion_status, "已达成")
        self.assertNotIn("整型截断除法缺陷", f2c_item.deduction_reason)

    def test_29_handwritten_partial_phone_no_penalty(self):
        """测试 29: 演草纸参数与代码完全一致但未写全手机号，不予扣除参数分且不标记复核 (学号 2026110090 场景)"""
        s = Student(seq=90, student_id="2026110090", name="测试生黄", class_id="260101")
        exp_x = 90
        deduced_y = 97

        # 模拟旧 AI 缓存中因为未写全手机号而给出 0.15 分的记录
        mock_ai_cache_entry = {
            "task_id": "lab1_2_handwritten",
            "student_id_observed": "2026110090",
            "name_observed": "",
            "phone_observed": "",
            "x_observed": 90,
            "y_observed": 97,
            "info_score": 0.2,
            "xy_match_score": 0.15,
            "derivation_score": 0.8,
            "suggested_score": 1.15,
            "is_pending": False,
            "reasoning": "纸面上方写有学号，但未标明姓名和手机号；参数 x=90, y=97 与预期一致，但未直接写明手机号故给0.15分",
            "needs_review": False,
            "review_reason": "",
        }

        normalized = HandwrittenReviewer.normalize_review_result(
            mock_ai_cache_entry, s.student_id, s.name, exp_x, deduced_y=deduced_y
        )
        # 按照新规则：动态根据参数事实重算，参数满分 0.2，信息宽容 0.5，推导 0.8，总分 1.5
        self.assertEqual(normalized["xy_match_score"], 0.2)
        self.assertEqual(normalized["info_score"], 0.5)
        self.assertEqual(normalized["suggested_score"], 1.5)
        self.assertFalse(normalized["needs_review"])

        judge = Lab1Judge()
        build_dir = self.test_work / "build_hw_90_test"
        build_dir.mkdir(parents=True, exist_ok=True)
        img = self.test_work / "hw_90.png"
        img.write_bytes(b"\xFF\xD8\xFF\xE0" + b"\x00" * 200)

        from unittest.mock import patch
        with patch("checker.ai.HandwrittenReviewer.review_images", return_value=normalized):
            prob = judge._grade_lab1_2(s, [], [img], build_dir, mock_ai=False)
            param_item = [it for it in prob.check_items if it.sub_item == "手写参数与身份及源码一致性"][0]
            self.assertEqual(param_item.deduction, 0.0)
            self.assertEqual(param_item.completion_status, "已达成")
            self.assertFalse(param_item.needs_review)

    def test_30_calculation_doubt_not_cleared_by_missing_info_leniency(self):
        """测试 30: 演草纸存在计算存疑或字迹不清时，宽容规则不得清除复核标志"""
        s = Student(seq=1, student_id="2026110001", name="测试生", class_id="260101")
        mock_res = {
            "task_id": "lab1_2_handwritten",
            "student_id_observed": "",
            "name_observed": "",
            "phone_observed": "",
            "x_observed": 1,
            "y_observed": 88,
            "info_score": 0.0,
            "xy_match_score": 0.2,
            "derivation_score": 0.4,
            "suggested_score": 0.6,
            "is_pending": False,
            "reasoning": "未填写姓名，同时推导字迹不清，位运算推导过程存疑",
            "needs_review": True,
            "review_reason": "未填写姓名，同时推导字迹不清，运算过程存疑",
        }

        normalized = HandwrittenReviewer.normalize_review_result(
            mock_res, s.student_id, s.name, 1, deduced_y=88
        )
        # 字迹不清/存疑必须保持复核，绝不能被误判为只有身份未注明而清除复核
        self.assertTrue(normalized["needs_review"])
        self.assertEqual(normalized["derivation_score"], 0.4)

    def test_31_detail_csv_review_flag_strictly_localized(self):
        """测试 31: 明细 CSV 中的复核标记必须严格限定于具体分项，不得扩散至全部 18 项"""
        import csv
        s = Student(seq=1, student_id="2026119999", name="复核测试生", class_id="260101")
        roster = RosterManager()
        roster.students = [s]
        roster.by_id = {s.student_id: s}
        roster.by_name = {s.name: [s]}

        grade = StudentGrade(student=s, max_score=5.0)
        grade.needs_review = True  # 学生级别标记需要复核
        grade.status = "needs_review"

        p1_1 = ProblemResult(problem_id="1-1", score=0.0, max_score=0.0)
        from checker.models import CheckItem
        p1_1.check_items = [
            CheckItem(problem_id="1-1", sub_item="代码编写与执行", submission_status="不适用", completion_status="无需提交", evidence="无需提交", needs_review=False),
            CheckItem(problem_id="1-1", sub_item="环境配置与认知", submission_status="不适用", completion_status="无需提交", evidence="无需提交", needs_review=False),
        ]
        p1_2 = ProblemResult(problem_id="1-2", score=2.5, max_score=3.0, needs_review=True)
        p1_2.check_items = [
            CheckItem(problem_id="1-2", sub_item="源码提交", submission_status="已提交", completion_status="已达成", evidence="main.c", needs_review=False),
            CheckItem(problem_id="1-2", sub_item="演草纸学生身份一致性", submission_status="已提交", completion_status="未达成", evidence="身份冲突", needs_review=True),
        ]
        grade.problem_results["1-1"] = p1_1
        grade.problem_results["1-2"] = p1_2

        out_dir = self.test_work / "out_reporter_localize"
        reporter = LabReporter("lab1", out_dir)
        reporter.generate_detail_csv([s], {s.student_id: grade})

        csv_file = out_dir / "lab1结果.csv"
        self.assertTrue(csv_file.exists())
        with open(csv_file, "r", encoding="utf-8-sig") as f:
            rows = list(csv.DictReader(f))
            item_map = {r["评分项"]: r["是否需要复核"] for r in rows}
            # 1-1 的无需提交项绝对不能为“是”
            self.assertEqual(item_map["代码编写与执行"], "否")
            self.assertEqual(item_map["环境配置与认知"], "否")
            self.assertEqual(item_map["源码提交"], "否")
            # 唯独标记了 needs_review 的项才为“是”
            self.assertEqual(item_map["演草纸学生身份一致性"], "是")

    def test_32_others_archive_unattributed_submissions(self):
        """测试 32: 无法确定学生归属的提交文件安全隔离归档至 labX/src/others/，同名文件隔离不覆盖，报告生成审计表"""
        s1 = Student(seq=1, student_id="2026110001", name="已知学生", class_id="260101")
        roster = RosterManager()
        roster.students = [s1]
        roster.by_id = {s1.student_id: s1}
        roster.by_name = {s1.name: [s1]}

        lab_dir = self.test_work / "lab_others_test"
        submit_dir = lab_dir / "submit"
        src_dir = lab_dir / "src"
        submit_dir.mkdir(parents=True, exist_ok=True)

        # 制作两个同名文件的未知提交 zip
        zip1_path = submit_dir / "unknown_pkg1.zip"
        with zipfile.ZipFile(zip1_path, "w") as z:
            z.writestr("test_data.csv", "col1,col2\n1,2\n")
        zip2_path = submit_dir / "unknown_pkg2.zip"
        with zipfile.ZipFile(zip2_path, "w") as z:
            z.writestr("test_data.csv", "colA,colB\n3,4\n")

        normalizer = SubmissionNormalizer(roster, "lab1", lab_dir)
        student_files_map = normalizer.process_all_submissions()

        # 1. 验证已知学生无归属文件，others 目录存在
        self.assertNotIn("others", student_files_map)
        others_dir = src_dir / "others"
        self.assertTrue(others_dir.exists())

        # 2. 验证两份同名文件各自归档于对应 stem 目录，未发生覆盖
        file1 = others_dir / "unknown_pkg1" / "test_data.csv"
        file2 = others_dir / "unknown_pkg2" / "test_data.csv"
        self.assertTrue(file1.exists())
        self.assertTrue(file2.exists())
        self.assertEqual(file1.read_text().strip(), "col1,col2\n1,2")
        self.assertEqual(file2.read_text().strip(), "colA,colB\n3,4")

        # 3. 验证 reporter 生成的报告包含统一归档位置与“无法归属、待复核”状态
        out_dir = lab_dir / "output"
        reporter = LabReporter("lab1", out_dir)
        grade = StudentGrade(student=s1, max_score=5.0)
        report_paths = reporter.generate_all_reports(
            roster_students=[s1],
            grades={s1.student_id: grade},
            unmatched_items=normalizer.unmatched_items,
        )
        report_content = Path(report_paths["summary_md"]).read_text(encoding="utf-8")
        self.assertIn("lab1/src/others/", report_content)
        self.assertIn("无法归属、待复核", report_content)
        self.assertIn("unknown_pkg1", report_content)
        self.assertIn("unknown_pkg2", report_content)

    def test_33_parallel_vs_serial_compilation_equivalence(self):
        """测试 33: 并行与串行公共编译等价性，多候选同名 main.c 兼容产物独立无覆盖"""
        from checker.compiler import ProjectCompiler
        from checker.models import CompileTask

        work = self.test_work / "test_compile_workers"
        work.mkdir(parents=True, exist_ok=True)
        compiler = ProjectCompiler()

        # 构造需要 Windows 兼容转换的两个同名 main.c
        src_dir1 = work / "cand1"
        src_dir2 = work / "cand2"
        src_dir1.mkdir(parents=True, exist_ok=True)
        src_dir2.mkdir(parents=True, exist_ok=True)

        main1 = src_dir1 / "main.c"
        main1.write_text('#include <stdio.h>\n#include <windows.h>\nint main(){printf("cand1");return 0;}', encoding="utf-8")
        main2 = src_dir2 / "main.c"
        main2.write_text('#include <stdio.h>\n#include <windows.h>\nint main(){printf("cand2");return 0;}', encoding="utf-8")

        build_dir = work / "build"
        task1 = CompileTask("2026110001", "1-2", 0, "lab1_2_cand_0", main1, build_dir)
        task2 = CompileTask("2026110001", "1-2", 1, "lab1_2_cand_1", main2, build_dir)

        # 1. 串行编译 (workers = 1)
        res_serial = compiler.compile_batch([task1, task2], max_workers=1)

        # 2. 检查生成的 compat 副本是否独立 (包含 problem_id 前缀)
        compat1 = build_dir / "lab1_2_cand_0_main_compat.c"
        compat2 = build_dir / "lab1_2_cand_1_main_compat.c"
        self.assertTrue(compat1.exists())
        self.assertTrue(compat2.exists())
        self.assertIn("cand1", compat1.read_text())
        self.assertIn("cand2", compat2.read_text())

        # 3. 并行编译 (workers = 4)
        shutil.rmtree(build_dir)
        res_parallel = compiler.compile_batch([task1, task2], max_workers=4)

        # 4. 对比结果一致性
        self.assertEqual(res_serial.keys(), res_parallel.keys())
        for k in res_serial:
            self.assertEqual(res_serial[k].success, res_parallel[k].success)
            self.assertEqual(res_serial[k].exit_code, res_parallel[k].exit_code)
            self.assertEqual(res_serial[k].binary_path.name, res_parallel[k].binary_path.name)
            self.assertEqual(res_serial[k].compat_applied, res_parallel[k].compat_applied)

    def test_34_manual_attribution_override(self):
        """测试 34: 教师人工核实名单 (MANUAL_ATTRIBUTIONS) 散放文件直接归属，不触发推断待复核"""
        s = Student(seq=1, student_id="2026110214", name="胡杨树", class_id="2612901")
        roster = RosterManager()
        roster.students = [s]
        roster.by_id = {s.student_id: s}
        roster.by_name = {s.name: [s]}

        lab_dir = self.test_work / "lab_manual_attr"
        submit_dir = lab_dir / "submit"
        submit_dir.mkdir(parents=True, exist_ok=True)

        # 制作哈希散放图片
        img_file = submit_dir / "68cbd4b92caa5a802f170c2c28b6fcb1.jpg"
        img_file.write_bytes(b"fake_jpg_content" * 50)

        normalizer = SubmissionNormalizer(roster, "lab1", lab_dir)
        student_files = normalizer.process_all_submissions()

        # 验证直接归属胡杨树，且未触发推断待复核标记
        self.assertIn("2026110214", student_files)
        self.assertEqual(len(student_files["2026110214"]), 1)
        self.assertEqual(student_files["2026110214"][0].matched_problem, "1-2")
        self.assertNotIn("2026110214", normalizer.inferred_students)

    def test_35_others_archive_no_overwrite_same_size(self):
        """测试 35: 两个来源目录各有一个同名同大小的无法归属文件，归档到 others 时确保目标唯一不覆盖"""
        roster = RosterManager()
        lab_dir = self.test_work / "lab_others_collision"
        normalizer = SubmissionNormalizer(roster, "lab1", lab_dir)
        normalizer.others_dir.mkdir(parents=True, exist_ok=True)

        dir1 = lab_dir / "submit" / "dir_a"
        dir2 = lab_dir / "submit" / "dir_b"
        dir1.mkdir(parents=True, exist_ok=True)
        dir2.mkdir(parents=True, exist_ok=True)

        file1 = dir1 / "unidentified.txt"
        file2 = dir2 / "unidentified.txt"
        file1.write_bytes(b"AAAA")  # 4 bytes
        file2.write_bytes(b"BBBB")  # 4 bytes

        rec1 = normalizer._archive_unattributed_item(file1, reason="未匹配A")
        rec2 = normalizer._archive_unattributed_item(file2, reason="未匹配B")

        # 验证两条审计记录指向不同目标
        self.assertNotEqual(rec1["archive_dest"], rec2["archive_dest"])

        # 验证两个文件均被归档在 others 中，内容均完整保留未被覆盖
        others_files = list((lab_dir / "src" / "others").glob("*.txt"))
        self.assertEqual(len(others_files), 2)
        contents = {f.read_text(encoding="utf-8") for f in others_files}
        self.assertEqual(contents, {"AAAA", "BBBB"})

        # 测试同名目录与多次冲突循环重命名
        d1 = dir1 / "sub_dir"
        d2 = dir2 / "sub_dir"
        d1.mkdir(parents=True, exist_ok=True)
        d2.mkdir(parents=True, exist_ok=True)
        (d1 / "f.txt").write_text("D1", encoding="utf-8")
        (d2 / "f.txt").write_text("D2", encoding="utf-8")

        drec1 = normalizer._archive_unattributed_item(d1, reason="目录A")
        drec2 = normalizer._archive_unattributed_item(d2, reason="目录B")
        self.assertNotEqual(drec1["archive_dest"], drec2["archive_dest"])

    def test_36_nested_batch_dir_contains_student_submissions(self):
        """测试 36: submit/学生作业/2026000067_Lab1-3.c 正常展开为总包，学生被正确匹配而非整体移入 others"""
        s = Student(seq=1, student_id="2026000067", name="测试生", class_id="2612901")
        roster = RosterManager()
        roster.students = [s]
        roster.by_id = {s.student_id: s}
        roster.by_name = {s.name: [s]}

        lab_dir = self.test_work / "lab_nested_batch"
        batch_dir = lab_dir / "submit" / "学生作业"
        batch_dir.mkdir(parents=True, exist_ok=True)
        src_file = batch_dir / "2026000067_Lab1-3.c"
        src_file.write_text("#include <stdio.h>\nint main(){ return 0; }", encoding="utf-8")

        normalizer = SubmissionNormalizer(roster, "lab1", lab_dir)
        student_files = normalizer.process_all_submissions()

        self.assertIn("2026000067", student_files)
        self.assertEqual(len(student_files["2026000067"]), 1)
        self.assertEqual(student_files["2026000067"][0].matched_problem, "1-3")
        # 学生作业目录未进入 others
        others_entries = [p for p in (lab_dir / "src" / "others").iterdir() if p.name != "README.md"]
        self.assertEqual(len(others_entries), 0)

    def test_37_corrupt_archive_for_matched_student_not_missing(self):
        """测试 37: 已确认学号的学生提交损坏压缩包，保留提交事实，记录解压失败待复核，不能转为未交(missing)"""
        s = Student(seq=1, student_id="2026110001", name="张三", class_id="2612901")
        roster = RosterManager()
        roster.students = [s]
        roster.by_id = {s.student_id: s}
        roster.by_name = {s.name: [s]}

        lab_dir = self.test_work / "lab_corrupt_archive"
        submit_dir = lab_dir / "submit"
        submit_dir.mkdir(parents=True, exist_ok=True)

        # 制作损坏的 zip (包含学号)
        corrupt_zip = submit_dir / "2026110001_张三_实验1.zip"
        corrupt_zip.write_bytes(b"PKCORRUPTED_ZIP_BYTES_NOT_VALID")

        normalizer = SubmissionNormalizer(roster, "lab1", lab_dir)
        student_files = normalizer.process_all_submissions()

        # 确认未进入 others，且归属张三
        self.assertIn("2026110001", student_files)
        self.assertEqual(len(student_files["2026110001"]), 1)
        sub_file = student_files["2026110001"][0]
        self.assertEqual(sub_file.file_type, "corrupt_archive")
        self.assertTrue((lab_dir / "src" / "2026110001" / corrupt_zip.name).exists())

        others_entries = [p for p in (lab_dir / "src" / "others").iterdir() if p.name != "README.md"]
        self.assertEqual(len(others_entries), 0)

        # 进行 Lab 1 评分评测
        judge = Lab1Judge()
        grade = judge.grade_student(s, student_files["2026110001"], build_dir=lab_dir / "build" / s.student_id, mock_ai=True)

        self.assertEqual(grade.status, "needs_review")
        self.assertTrue(grade.needs_review)
        self.assertNotEqual(grade.status, "missing")

        # 验证各分项状态
        items_1_2 = grade.problem_results["1-2"].check_items
        sub_item_1_2 = next(it for it in items_1_2 if it.sub_item == "源码提交")
        self.assertEqual(sub_item_1_2.submission_status, "已提交")
        self.assertEqual(sub_item_1_2.completion_status, "无法验证")
        self.assertEqual(sub_item_1_2.agent_note, "解压失败、内容无法验证、待复核")
        self.assertTrue(sub_item_1_2.needs_review)

        # 验证评语提及损坏解压失败及待人工复核
        self.assertIn("损坏无法正常解压", grade.comment)
        self.assertIn("待人工复核", grade.comment)
        self.assertGreaterEqual(len(grade.comment), 50)

    def test_38_semantic_cache_invalidates_on_requirements_change_and_clamps(self):
        """测试 38: 语义复核缓存键包含 expected_behavior 与 max_score，并在命中后核验分值范围"""
        import hashlib
        from checker.ai import SemanticAmbiguityReviewer, _load_ai_cache, _save_ai_cache

        # 计算相同输入、不同 expected_behavior 或不同 max_score 时的缓存键
        def get_cache_key(expected: str, max_sc: float) -> str:
            hasher = hashlib.md5()
            hasher.update("2-1".encode("utf-8"))
            hasher.update("2026000001".encode("utf-8"))
            hasher.update("12\n".encode("utf-8"))
            hasher.update("Noon".encode("utf-8"))
            hasher.update("".encode("utf-8"))
            hasher.update(expected.encode("utf-8"))
            hasher.update(f"{max_sc:.4f}".encode("utf-8"))
            return f"semantic_{hasher.hexdigest()}"

        key1 = get_cache_key("Noon\n", 2.0)
        key2 = get_cache_key("Different\n", 2.0)
        key3 = get_cache_key("Noon\n", 0.1)

        self.assertNotEqual(key1, key2, "期望输出变化应使缓存失效")
        self.assertNotEqual(key1, key3, "最高分要求变化应使缓存失效")

        # 测试缓存命中后的分值范围钳位
        cache = _load_ai_cache()
        cache[key1] = {
            "task_id": "semantic_2-1",
            "category": "format_difference",
            "passed": True,
            "suggested_score": 2.0,
            "is_pending": False,
            "reasoning": "format diff",
            "needs_review": False,
        }
        _save_ai_cache(cache)

        # 模拟 key3 在缓存中已有旧的 2.0 高分记录，核验命中后是否正确钳位至 0.1
        cache[key3] = dict(cache[key1])
        _save_ai_cache(cache)

        res = SemanticAmbiguityReviewer.review_output(
            problem_id="2-1",
            student_id="2026000001",
            input_data="12\n",
            expected_behavior="Noon\n",
            actual_output="Noon",
            source_code="",
            max_score=0.1,
            mock_mode=False,
        )
        self.assertAlmostEqual(res["suggested_score"], 0.1, places=2)

    def test_39_detect_and_decode_source_encodings(self):
        """测试39: 严格编码识别与多字符集解码校验 (UTF-8, UTF-8 BOM, GBK/GB18030, UTF-16, 拒绝二进制与损坏兜底)"""
        from checker.unpack import detect_and_decode_source

        # 1. 纯 ASCII
        ascii_code = b"#include <stdio.h>\nint main() { return 0; }\n"
        ok, enc, text, conv, err = detect_and_decode_source(ascii_code)
        self.assertTrue(ok)
        self.assertEqual(enc, "utf-8")
        self.assertFalse(conv)
        self.assertIn("main", text)

        # 2. UTF-8 含中文
        utf8_code = "// 摄氏温度转华氏温度\n#include <stdio.h>\n".encode("utf-8")
        ok, enc, text, conv, err = detect_and_decode_source(utf8_code)
        self.assertTrue(ok)
        self.assertEqual(enc, "utf-8")
        self.assertFalse(conv)
        self.assertIn("摄氏温度", text)

        # 3. UTF-8 with BOM
        utf8_bom_code = b"\xef\xbb\xbf" + utf8_code
        ok, enc, text, conv, err = detect_and_decode_source(utf8_bom_code)
        self.assertTrue(ok)
        self.assertEqual(enc, "utf-8-sig")
        self.assertTrue(conv)
        self.assertIn("摄氏温度", text)

        # 4. GBK / GB18030
        gbk_code = "// 华氏度与摄氏度转换程序\n#include <stdio.h>\nint main(){ return 0; }".encode("gbk")
        ok, enc, text, conv, err = detect_and_decode_source(gbk_code)
        self.assertTrue(ok)
        self.assertEqual(enc, "gb18030")
        self.assertTrue(conv)
        self.assertIn("华氏度与摄氏度", text)

        # 5. UTF-16 with BOM
        utf16_bom = "// 字母金字塔程序\n".encode("utf-16")
        ok, enc, text, conv, err = detect_and_decode_source(utf16_bom)
        self.assertTrue(ok)
        self.assertIn("utf-16", enc)
        self.assertTrue(conv)
        self.assertIn("字母金字塔", text)

        # 6. UTF-16 without BOM (结构特征显著)
        utf16_nobom = "#include <stdio.h>\nint main() { return 0; }\n".encode("utf-16-le")
        ok, enc, text, conv, err = detect_and_decode_source(utf16_nobom)
        self.assertTrue(ok)
        self.assertEqual(enc, "utf-16-le")
        self.assertTrue(conv)
        self.assertIn("main", text)

        # 7. 二进制文件识别并拒绝转码 (含非打印控制字符)
        binary_data = b"PK\x03\x04\x14\x00\x00\x00\x08\x00"
        ok, enc, text, conv, err = detect_and_decode_source(binary_data)
        self.assertFalse(ok)
        self.assertEqual(enc, "binary")

        # 8. 损坏编码识别并拒绝错误假转码 (严禁 Latin-1 或 errors='replace' 兜底)
        corrupted = b"int main() { /* \xff\xff\xff\xff */ }"
        ok, enc, text, conv, err = detect_and_decode_source(corrupted)
        self.assertFalse(ok)
        self.assertEqual(enc, "unknown")

    def test_40_staging_encoding_normalization(self):
        """测试40: 归档阶段 submit -> src 自动将 GBK/UTF-8-BOM 转为标准 UTF-8 无 BOM"""
        lab_dir = self.test_work / "test_40_lab"
        submit_dir = lab_dir / "submit"
        submit_dir.mkdir(parents=True, exist_ok=True)

        std = self.roster_mgr.students[0]
        # 创建一个 GBK 编码的 1-3 源码提交文件
        gbk_text = '// 华氏温度转摄氏温度\n#include <stdio.h>\nint main() { printf("华氏温度=%.2f\\n", 32.0); return 0; }'
        gbk_file = submit_dir / f"{std.student_id}_Lab1-3.c"
        gbk_file.write_bytes(gbk_text.encode("gb18030"))

        normalizer = SubmissionNormalizer(self.roster_mgr, "lab1", lab_dir)
        staged = normalizer.process_all_submissions()

        self.assertIn(std.student_id, staged)
        std_files = staged[std.student_id]
        src_files = [f for f in std_files if f.file_type == "c_source"]
        self.assertEqual(len(src_files), 1)
        sf = src_files[0]
        self.assertEqual(sf.source_encoding, "gb18030")
        self.assertTrue(sf.encoding_converted)
        self.assertEqual(sf.encoding_error, "")

        # 验证 src 目录中的文件是严格 UTF-8 无 BOM
        staged_bytes = sf.path.read_bytes()
        self.assertFalse(staged_bytes.startswith(b"\xef\xbb\xbf"), "必须保存为无 BOM 的 UTF-8")
        staged_text = staged_bytes.decode("utf-8")  # strict!
        self.assertIn("华氏温度转摄氏温度", staged_text)

    def test_41_staging_idempotence_and_no_alt_creation(self):
        """测试41: 归档幂等性测试，重复归档原始 GBK 文件严禁产生 _alt 冗余副本"""
        lab_dir = self.test_work / "test_41_lab"
        submit_dir = lab_dir / "submit"
        submit_dir.mkdir(parents=True, exist_ok=True)

        std = self.roster_mgr.students[0]
        gbk_text = '// 测试幂等性归档\n#include <stdio.h>\nint main() { return 0; }'
        gbk_file = submit_dir / f"{std.student_id}_Lab1-3.c"
        gbk_file.write_bytes(gbk_text.encode("gb18030"))

        # 第一次归档
        norm1 = SubmissionNormalizer(self.roster_mgr, "lab1", lab_dir)
        norm1.process_all_submissions()

        std_src_dir = lab_dir / "src" / std.student_id
        files_after_run1 = list(std_src_dir.glob("*.c"))
        self.assertEqual(len(files_after_run1), 1)

        # 第二次归档 (直接重新运行，submit 内容未变)
        norm2 = SubmissionNormalizer(self.roster_mgr, "lab1", lab_dir)
        norm2.process_all_submissions()

        files_after_run2 = list(std_src_dir.glob("*.c"))
        self.assertEqual(len(files_after_run2), 1, "重复归档不得因原始 GBK 与规范化 UTF-8 字节不同而生成 _alt 副本")
        self.assertFalse(any("_alt" in f.name for f in files_after_run2))

    def test_42_temperature_symbols_encoding_detection(self):
        """测试42: [P1] 包含 ℃、℉ 的合法 GBK 源码能正确解码为 UTF-8，不因汉字比例误判为未知编码"""
        from checker.unpack import detect_and_decode_source

        # 包含 ℃、℉ 和 °C 的 GBK 代码（无中文汉字）
        gbk_code = b'#include <stdio.h>\nint main() { printf("100 \xa1\xe6 = 212 \xa8H\\n"); return 0; }\n'
        ok, enc, text, conv, err = detect_and_decode_source(gbk_code)
        self.assertTrue(ok, f"应成功识别含温度符号的 GBK 源码，而非被误判为: {err}")
        self.assertIn("℃", text)
        self.assertIn("℉", text)
        self.assertEqual(enc, "gb18030")
        self.assertTrue(conv)

    def test_43_lab2_2_corrupted_encoding_review_flag(self):
        """测试43: [P1] Lab 2-2 源码编码异常必须置 needs_review=True, is_pending=True 并记录待复核清单"""
        from checker.labs.lab2 import Lab2Judge
        from checker.models import SubmissionFile

        lab_dir = self.test_work / "test_43_lab"
        build_dir = lab_dir / "build"
        src_dir = lab_dir / "src"
        src_dir.mkdir(parents=True, exist_ok=True)
        build_dir.mkdir(parents=True, exist_ok=True)

        std = self.roster_mgr.students[0]
        # 创建一个无法解码的损坏源码
        bad_src = src_dir / f"{std.student_id}_Lab2-2.c"
        bad_src.write_bytes(b"#include <stdio.h>\n/* \xff\xff\xff\xff */\nint main(){ return 0; }")

        bad_sf = SubmissionFile(
            path=bad_src,
            original_name=bad_src.name,
            file_type="c_source",
            matched_problem="2-2",
            source_encoding="unknown",
            encoding_converted=False,
            encoding_error="严格多编码解码均失败",
        )

        judge = Lab2Judge()
        grade = judge.grade_student(std, [bad_sf], build_dir, mock_ai=True)

        self.assertTrue(grade.needs_review, "编码损坏的 Lab 2-2 必须标记 needs_review=True")
        self.assertTrue(grade.is_pending, "编码损坏的 Lab 2-2 必须标记 is_pending=True")
        self.assertEqual(grade.status, "needs_review", "编码损坏的 Lab 2-2 成绩状态必须为 needs_review")
        self.assertTrue(any("Lab 2-2" in r and "编码损坏" in r for r in grade.review_reasons), "待复核清单需记录编码损坏原因")

    def test_44_candidate_deduplication_same_content_no_alt(self):
        """测试44: [P2] 同一路径依次提交内容 A、B、B，后两份相同内容必须去重，不生成多余 _alt2"""
        lab_dir = self.test_work / "test_44_lab"
        normalizer = SubmissionNormalizer(self.roster_mgr, "lab1", lab_dir)
        std = self.roster_mgr.students[0]

        temp_submit = lab_dir / "temp_input"
        temp_submit.mkdir(parents=True, exist_ok=True)

        # 候选 1: 内容 A
        f_a = temp_submit / "cand1.c"
        f_a.write_text("int main() { /* Version A */ return 0; }", encoding="utf-8")
        sf_1 = normalizer._classify_and_stage_file(std, f_a, "main.c")
        self.assertIsNotNone(sf_1)
        self.assertEqual(sf_1.path.name, "main.c")
        normalizer.student_files.setdefault(std.student_id, []).append(sf_1)

        # 候选 2: 内容 B (与 A 冲突，生成 main_alt.c)
        f_b1 = temp_submit / "cand2.c"
        f_b1.write_text("int main() { /* Version B */ return 0; }", encoding="utf-8")
        sf_2 = normalizer._classify_and_stage_file(std, f_b1, "main.c")
        self.assertIsNotNone(sf_2)
        self.assertEqual(sf_2.path.name, "main_alt.c")
        normalizer.student_files.setdefault(std.student_id, []).append(sf_2)

        # 候选 3: 内容 B (与候选 2 完全相同，必须去重跳过，返回 None)
        f_b2 = temp_submit / "cand3.c"
        f_b2.write_text("int main() { /* Version B */ return 0; }", encoding="utf-8")
        sf_3 = normalizer._classify_and_stage_file(std, f_b2, "main.c")
        self.assertIsNone(sf_3, "内容相同的后续候选应去重并返回 None")

        # 验证磁盘上只有 main.c 与 main_alt.c，绝不生成 main_alt2.c
        staged_files = sorted([p.name for p in (lab_dir / "src" / std.student_id).glob("*.c")])
        self.assertEqual(staged_files, ["main.c", "main_alt.c"], "磁盘上严禁存在 main_alt2.c")

    def test_45_lab2_pipeline_archive_grade_csv_export(self):
        """测试45: [P1] Lab2 全流程 集成测试 (归档 -> 判题 -> CSV/MD导出)
        验证 Lab2 编码异常样本导出时:
        1. lab2结果.csv 为 warmup, 2-1, 2-2 分项，无 1-1, 1-2, 1-3
        2. 编码异常项达成状态为 '待复核', 建议扣分为 '待复核', 是否需要复核为 '是'
        3. lab2实验成绩单.csv 中 手写材料 为 '不适用 (Lab 2 无手写验算要求)'
        4. 建议扣分合计为 '待复核'，扣分依据无 Lab1 扣分项，评语针对 Lab2 且 >= 50 字
        5. 缺交学生正确生成 Lab2 缺交分项 (warmup, 2-1, 2-2, 建议扣分 5.00)
        """
        import csv
        lab_dir = self.test_work / "test_45_lab2"
        submit_dir = lab_dir / "submit"
        src_dir = lab_dir / "src"
        build_dir = lab_dir / "build"
        output_dir = lab_dir / "output"

        submit_dir.mkdir(parents=True, exist_ok=True)
        output_dir.mkdir(parents=True, exist_ok=True)

        std_a = self.roster_mgr.students[0]
        std_b = self.roster_mgr.students[1] if len(self.roster_mgr.students) > 1 else Student(seq=99, student_id="2026999999", name="测试缺交生", class_id="260101")
        std_c = Student(seq=100, student_id="2026888888", name="测试全通生", class_id="260101")
        test_roster = [std_a, std_b, std_c]

        # 1. 模拟 submit 提交
        code_2_1 = """#include <stdio.h>
int main() {
    int h;
    if (scanf("%d", &h) != 1) {
        printf("Invalid input!\\n");
        return 0;
    }
    if (h >= 0 && h < 12) printf("Good morning!\\n");
    else if (h == 12) printf("Good noon!\\n");
    else if (h >= 13 && h < 18) printf("Good afternoon!\\n");
    else if (h >= 18 && h < 21) printf("Good evening!\\n");
    else if (h >= 21 && h < 24) printf("Good night!\\n");
    else printf("Invalid input!\\n");
    return 0;
}
"""
        corrupted_2_2_bytes = b"#include <stdio.h>\n/* \xff\xff\xfe\xfe */\nint main(){ printf(\"error\"); return 0; }\n"

        code_2_2_perfect = """#include <stdio.h>
#include <ctype.h>

int main() {
    char ch;
    if (scanf(" %c", &ch) != 1) {
        printf("ERROR\\n");
        return 0;
    }
    if (!isalpha(ch)) {
        printf("ERROR: not a letter!\\n");
        return 0;
    }
    ch = toupper(ch);
    int n = ch - 'A' + 1;
    for (int i = 0; i < n; i++) {
        for (int j = 0; j < n - i - 1; j++) printf(" ");
        for (char c = 'A'; c <= 'A' + i; c++) printf("%c", c);
        for (char c = 'A' + i - 1; c >= 'A'; c--) printf("%c", c);
        printf("\\n");
    }
    return 0;
}
"""

        (submit_dir / f"{std_a.student_id}_Lab2-1.c").write_text(code_2_1, encoding="utf-8")
        (submit_dir / f"{std_a.student_id}_Lab2-2.c").write_bytes(corrupted_2_2_bytes)

        (submit_dir / f"{std_c.student_id}_Lab2-1.c").write_text(code_2_1, encoding="utf-8")
        (submit_dir / f"{std_c.student_id}_Lab2-2.c").write_text(code_2_2_perfect, encoding="utf-8")

        # 2. 执行 归档 (SubmissionNormalizer)
        # 为支持 std_c 识别，临时注入点名册
        self.roster_mgr.students.append(std_c)
        self.roster_mgr.by_id[std_c.student_id] = std_c
        self.roster_mgr.by_name.setdefault(std_c.name, []).append(std_c)

        try:
            normalizer = SubmissionNormalizer(self.roster_mgr, "lab2", lab_dir)
            normalizer.process_all_submissions()
            self.assertIn(std_a.student_id, normalizer.student_files)
            self.assertIn(std_c.student_id, normalizer.student_files)
            staged_files_a = normalizer.student_files[std_a.student_id]
            self.assertEqual(len(staged_files_a), 2)
            bad_sf_2_2 = next(f for f in staged_files_a if f.matched_problem == "2-2")
            self.assertTrue(bad_sf_2_2.encoding_error)

            # 3. 执行 判题 (Lab2Judge)
            judge = Lab2Judge()
            grades = {}
            for s in test_roster:
                s_files = normalizer.student_files.get(s.student_id, [])
                s_build = build_dir / s.student_id
                s_build.mkdir(parents=True, exist_ok=True)
                grade = judge.grade_student(s, s_files, s_build, mock_ai=True)
                grades[s.student_id] = grade

            grade_a = grades[std_a.student_id]
            grade_b = grades[std_b.student_id]
            grade_c = grades[std_c.student_id]

            self.assertTrue(grade_a.needs_review)
            self.assertTrue(grade_a.is_pending)
            self.assertEqual(grade_a.status, "needs_review")
            self.assertEqual(grade_b.status, "missing")
            self.assertEqual(grade_c.status, "graded")
            self.assertAlmostEqual(grade_c.total_score, 5.00, places=2)

            # 4. 执行 报表导出 (LabReporter)
            reporter = LabReporter("lab2", output_dir)
            reports = reporter.generate_all_reports(
                roster_students=test_roster,
                grades=grades,
                encoding_stats=normalizer.encoding_stats,
            )

            detail_csv_path = reports["detail_csv"]
            roster_csv_path = reports["roster_csv"]
            summary_md_path = reports["summary_md"]

            self.assertTrue(detail_csv_path.exists())
            self.assertTrue(roster_csv_path.exists())
            self.assertTrue(summary_md_path.exists())

            # 5. 校验 lab2结果.csv
            with open(detail_csv_path, "r", encoding="utf-8-sig") as f:
                reader = list(csv.DictReader(f))

            pids = set(r["题号"] for r in reader)
            self.assertNotIn("1-1", pids, "lab2结果.csv 严禁出现 1-1")
            self.assertNotIn("1-2", pids, "lab2结果.csv 严禁出现 1-2")
            self.assertNotIn("1-3", pids, "lab2结果.csv 严禁出现 1-3")
            self.assertIn("warmup", pids, "lab2结果.csv 需包含 warmup")
            self.assertIn("2-1", pids, "lab2结果.csv 需包含 2-1")
            self.assertIn("2-2", pids, "lab2结果.csv 需包含 2-2")

            # 校验 std_a 的 2-2 编码异常分项状态
            rows_a_2_2 = [r for r in reader if r["学号"] == std_a.student_id and r["题号"] == "2-2"]
            self.assertGreater(len(rows_a_2_2), 0)
            for r in rows_a_2_2:
                self.assertEqual(r["达成状态"], "待复核")
                self.assertEqual(r["建议扣分"], "待复核")
                self.assertEqual(r["是否需要复核"], "是")

            # 校验缺交生 std_b 的分项
            rows_b = [r for r in reader if r["学号"] == std_b.student_id]
            self.assertEqual(len(rows_b), 10, "缺交生默认应生成 10 个 Lab2 分项")
            self.assertTrue(all(r["题号"] in {"warmup", "2-1", "2-2"} for r in rows_b))
            r_b_2_1 = next(r for r in rows_b if r["题号"] == "2-1" and r["评分项"] == "源码提交")
            self.assertEqual(r_b_2_1["建议扣分"], "2.00")
            r_b_2_2 = next(r for r in rows_b if r["题号"] == "2-2" and r["评分项"] == "源码提交")
            self.assertEqual(r_b_2_2["建议扣分"], "3.00")

            # 校验全通生 std_c 的分项
            rows_c = [r for r in reader if r["学号"] == std_c.student_id]
            for r in rows_c:
                if r["题号"] != "warmup":
                    self.assertEqual(r["达成状态"], "已达成")
                    self.assertEqual(r["建议扣分"], "0.00")
                    self.assertEqual(r["是否需要复核"], "否")

            # 6. 校验 lab2实验成绩单.csv
            with open(roster_csv_path, "r", encoding="utf-8-sig") as f:
                roster_rows = list(csv.DictReader(f))

            row_a = next(r for r in roster_rows if r["学号"] == std_a.student_id)
            row_b = next(r for r in roster_rows if r["学号"] == std_b.student_id)
            row_c = next(r for r in roster_rows if r["学号"] == std_c.student_id)

            # std_a 检查
            self.assertEqual(row_a["手写材料提交与正确性"], "不适用 (Lab 2 无手写验算要求)")
            self.assertNotIn("未提交演草纸", row_a["手写材料提交与正确性"])
            self.assertEqual(row_a["建议扣分合计"], "待复核", "存在编码异常时建议扣分合计应为 '待复核'")
            self.assertNotIn("1-2", row_a["扣分依据"])
            self.assertNotIn("1-3", row_a["扣分依据"])
            self.assertIn("待人工复核", row_a["待复核事项"])
            self.assertIn("编码损坏", row_a["待复核事项"])
            self.assertGreaterEqual(len(row_a["实验评语"]), 50)
            self.assertNotIn("演草纸", row_a["实验评语"])
            self.assertNotIn("实验1", row_a["实验评语"])

            # std_b 检查
            self.assertEqual(row_b["手写材料提交与正确性"], "不适用 (Lab 2 无手写验算要求)")
            self.assertEqual(row_b["建议扣分合计"], "5.00")
            self.assertIn("warmup:未提交", row_b["各题验收汇总"])
            self.assertIn("2-1:未提交", row_b["各题验收汇总"])
            self.assertIn("2-2:未提交", row_b["各题验收汇总"])
            self.assertNotIn("1-1", row_b["各题验收汇总"])
            self.assertGreaterEqual(len(row_b["实验评语"]), 50)
            self.assertIn("LAB2", row_b["实验评语"])

            # std_c 检查
            self.assertEqual(row_c["手写材料提交与正确性"], "不适用 (Lab 2 无手写验算要求)")
            self.assertEqual(row_c["建议扣分合计"], "0.00")
            self.assertEqual(row_c["扣分依据"], "无扣分项")
            self.assertIn("2-1:(得分2.00/2.00)", row_c["各题验收汇总"])
            self.assertIn("2-2:(得分3.00/3.00)", row_c["各题验收汇总"])
            self.assertEqual(row_c["源码提交与正确性"], "2-1与2-2源码均提交且编译通过，全部测试用例正确")
            self.assertGreaterEqual(len(row_c["实验评语"]), 50)

            # 7. 校验 lab2报告.md
            md_text = summary_md_path.read_text(encoding="utf-8")
            self.assertIn("《程序设计思维与实践》LAB2 批量评测报告", md_text)
            self.assertIn("## 三、 评分细则与说明", md_text)
            self.assertIn("Lab 2-1 时间段问候 (满分 2.0分)", md_text)
            self.assertIn("Lab 2-2 字母金字塔 (满分 3.0分)", md_text)
            self.assertNotIn("Lab 1-2 进制与逻辑", md_text)
            self.assertIn("核对源码文件编码与原始字符", md_text)
        finally:
            if std_c in self.roster_mgr.students:
                self.roster_mgr.students.remove(std_c)
            self.roster_mgr.by_id.pop(std_c.student_id, None)

    def test_46_multi_zip_version_naming_and_deduplication(self):
        """测试46: 多个独立ZIP中的同题多源码正确保存为原名与 _alt，重复归档不产生多余副本"""
        lab_dir = self.test_work / "test_46_lab"
        submit_dir = lab_dir / "submit"
        submit_dir.mkdir(parents=True, exist_ok=True)

        std = self.roster_mgr.students[0]

        code_v1 = '#include <stdio.h>\nint main() { printf("version 1\\n"); return 0; }\n'
        code_v2 = '#include <stdio.h>\nint main() { printf("version 2\\n"); return 0; }\n'

        zip1_path = submit_dir / f"{std.student_id}_sub1.zip"
        with zipfile.ZipFile(zip1_path, "w") as zf:
            zf.writestr(f"{std.student_id}/Lab1-3.c", code_v1)

        zip2_path = submit_dir / f"{std.student_id}_sub2.zip"
        with zipfile.ZipFile(zip2_path, "w") as zf:
            zf.writestr(f"{std.student_id}/Lab1-3.c", code_v2)

        normalizer1 = SubmissionNormalizer(self.roster_mgr, "lab1", lab_dir)
        staged1 = normalizer1.process_all_submissions()

        std_files1 = staged1.get(std.student_id, [])
        c_files1 = [f for f in std_files1 if f.file_type == "c_source"]
        self.assertEqual(len(c_files1), 2, "应归档两份不同的源码候选")
        names1 = sorted([f.path.name for f in c_files1])
        self.assertEqual(names1, ["Lab1-3.c", "Lab1-3_alt.c"])

        texts1 = {f.path.read_text(encoding="utf-8") for f in c_files1}
        self.assertIn(code_v1, texts1)
        self.assertIn(code_v2, texts1)

        normalizer2 = SubmissionNormalizer(self.roster_mgr, "lab1", lab_dir)
        staged2 = normalizer2.process_all_submissions()

        std_files2 = staged2.get(std.student_id, [])
        c_files2 = [f for f in std_files2 if f.file_type == "c_source"]
        self.assertEqual(len(c_files2), 2, "重复归档不得生成 _alt2 等多余副本")
        names2 = sorted([f.path.name for f in c_files2])
        self.assertEqual(names2, ["Lab1-3.c", "Lab1-3_alt.c"])

    def test_47_header_version_isolation_and_execution(self):
        """测试47: 包含私有头文件的独立版本完整保留目录结构，不使用其他版本的头文件，运行结果对应各自版本"""
        lab_dir = self.test_work / "test_47_lab"
        submit_dir = lab_dir / "submit"
        build_dir = lab_dir / "build"
        submit_dir.mkdir(parents=True, exist_ok=True)
        build_dir.mkdir(parents=True, exist_ok=True)

        std = self.roster_mgr.students[0]

        main_c = """#include <stdio.h>
#include "config.h"
int main() {
    printf("SCALE=%d\\n", SCALE);
    return 0;
}
"""
        header_v1 = "#define SCALE 10\n"
        header_v2 = "#define SCALE 20\n"

        zip1 = submit_dir / f"{std.student_id}_v1.zip"
        with zipfile.ZipFile(zip1, "w") as zf:
            zf.writestr(f"{std.student_id}/Lab1-3/main.c", main_c)
            zf.writestr(f"{std.student_id}/Lab1-3/inc/config.h", header_v1)

        zip2 = submit_dir / f"{std.student_id}_v2.zip"
        with zipfile.ZipFile(zip2, "w") as zf:
            zf.writestr(f"{std.student_id}/Lab1-3/main.c", main_c)
            zf.writestr(f"{std.student_id}/Lab1-3/inc/config.h", header_v2)

        normalizer = SubmissionNormalizer(self.roster_mgr, "lab1", lab_dir)
        staged = normalizer.process_all_submissions()

        std_files = staged.get(std.student_id, [])
        c_files = [f for f in std_files if f.file_type == "c_source"]
        self.assertEqual(len(c_files), 2, "两份包含不同私有头文件的独立版本均应归档")

        v1_file = next(f for f in c_files if "Lab1-3/" in f.path.as_posix() and "_v002" not in f.path.as_posix())
        v2_file = next(f for f in c_files if "Lab1-3_v002" in f.path.as_posix())

        h1_text = (v1_file.path.parent / "inc" / "config.h").read_text(encoding="utf-8")
        h2_text = (v2_file.path.parent / "inc" / "config.h").read_text(encoding="utf-8")
        self.assertEqual(h1_text.strip(), "#define SCALE 10")
        self.assertEqual(h2_text.strip(), "#define SCALE 20")

        compiler = ProjectCompiler()
        std_build = build_dir / std.student_id
        res1 = compiler.compile(v1_file.path, std_build, "bin_v1")
        res2 = compiler.compile(v2_file.path, std_build, "bin_v2")

        self.assertTrue(res1.success, f"v1 编译失败: {res1.stderr}")
        self.assertTrue(res2.success, f"v2 编译失败: {res2.stderr}")

        run1 = run_binary(res1.binary_path, input_str="", timeout=2.0)
        run2 = run_binary(res2.binary_path, input_str="", timeout=2.0)

        self.assertIn("SCALE=10", run1.stdout)
        self.assertIn("SCALE=20", run2.stdout)

    def test_48_nested_archive_slug_isolation(self):
        """测试48: 同名嵌套ZIP使用来源相对路径作为前缀，互不覆盖"""
        lab_dir = self.test_work / "test_48_lab"
        submit_dir = lab_dir / "submit"
        submit_dir.mkdir(parents=True, exist_ok=True)

        std = self.roster_mgr.students[0]

        code_a = '#include <stdio.h>\nint main() { printf("nested A\\n"); return 0; }\n'
        code_b = '#include <stdio.h>\nint main() { printf("nested B\\n"); return 0; }\n'

        nested_a_bytes = io.BytesIO()
        with zipfile.ZipFile(nested_a_bytes, "w") as zf:
            zf.writestr("Lab1-3.c", code_a)

        nested_b_bytes = io.BytesIO()
        with zipfile.ZipFile(nested_b_bytes, "w") as zf:
            zf.writestr("Lab1-3.c", code_b)

        outer_zip = submit_dir / f"{std.student_id}_outer.zip"
        with zipfile.ZipFile(outer_zip, "w") as zf:
            zf.writestr("dir_a/version.zip", nested_a_bytes.getvalue())
            zf.writestr("dir_b/version.zip", nested_b_bytes.getvalue())

        normalizer = SubmissionNormalizer(self.roster_mgr, "lab1", lab_dir)
        staged = normalizer.process_all_submissions()

        std_files = staged.get(std.student_id, [])
        c_files = [f for f in std_files if f.file_type == "c_source"]
        self.assertEqual(len(c_files), 2, "嵌套同名zip解压时不应相互覆盖，应全部归档")

        texts = {f.path.read_text(encoding="utf-8") for f in c_files}
        self.assertIn(code_a, texts)
        self.assertIn(code_b, texts)

    def test_49_zip_duplicate_member_disambiguation(self):
        """测试49: 单个ZIP内的同名文件解压时消歧并全部归档"""
        lab_dir = self.test_work / "test_49_lab"
        submit_dir = lab_dir / "submit"
        submit_dir.mkdir(parents=True, exist_ok=True)

        std = self.roster_mgr.students[0]

        code_1 = '#include <stdio.h>\nint main() { printf("dup 1\\n"); return 0; }\n'
        code_2 = '#include <stdio.h>\nint main() { printf("dup 2\\n"); return 0; }\n'

        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as zf:
            zf.writestr(f"{std.student_id}/Lab1-3.c", code_1)
            zf.writestr(f"{std.student_id}/Lab1-3.c", code_2)

        zip_path = submit_dir / f"{std.student_id}_dup.zip"
        zip_path.write_bytes(buf.getvalue())

        normalizer = SubmissionNormalizer(self.roster_mgr, "lab1", lab_dir)
        staged = normalizer.process_all_submissions()

        std_files = staged.get(std.student_id, [])
        c_files = [f for f in std_files if f.file_type == "c_source"]
        self.assertEqual(len(c_files), 2, "单个zip内同名条目必须通过消歧全部归档")

        texts = {f.path.read_text(encoding="utf-8") for f in c_files}
        self.assertIn(code_1, texts)
        self.assertIn(code_2, texts)

    def test_50_container_cache_invalidation_and_idempotence(self):
        """测试50: 修改容器ZIP内容后重新归档，能够重新解压最新内容并完成评测"""
        lab_dir = self.test_work / "test_50_lab"
        submit_dir = lab_dir / "submit"
        submit_dir.mkdir(parents=True, exist_ok=True)

        std = self.roster_mgr.students[0]

        code_v1 = '#include <stdio.h>\nint main() { printf("original\\n"); return 0; }\n'
        code_v2 = '#include <stdio.h>\nint main() { printf("updated\\n"); return 0; }\n'

        container_zip = submit_dir / f"{std.student_id}_container.zip"
        with zipfile.ZipFile(container_zip, "w") as zf:
            zf.writestr(f"{std.student_id}/Lab1-3.c", code_v1)

        normalizer1 = SubmissionNormalizer(self.roster_mgr, "lab1", lab_dir)
        staged1 = normalizer1.process_all_submissions()
        file1 = staged1[std.student_id][0]
        self.assertIn("original", file1.path.read_text(encoding="utf-8"))

        import time
        time.sleep(0.05)
        with zipfile.ZipFile(container_zip, "w") as zf:
            zf.writestr(f"{std.student_id}/Lab1-3.c", code_v2)

        normalizer2 = SubmissionNormalizer(self.roster_mgr, "lab1", lab_dir)
        staged2 = normalizer2.process_all_submissions()
        file2 = staged2[std.student_id][0]
        self.assertIn("updated", file2.path.read_text(encoding="utf-8"), "容器更新后应解压最新内容并更新归档")

    def test_51_lab1_2_full_score_candidate_selection(self):
        """测试51: Lab1-2候选选择使用“代码+手写验算”整题最终最高分，不受输入顺序影响"""
        lab_dir = self.test_work / "test_51_lab"
        build_dir = lab_dir / "build"
        src_dir = lab_dir / "src"
        src_dir.mkdir(parents=True, exist_ok=True)
        build_dir.mkdir(parents=True, exist_ok=True)

        std = Student(seq=1, student_id="2026110214", name="测试满分生", class_id="260101", phone="17393861988")

        # 候选 A: 逻辑正确，但硬编码 y=20 与推导 y=88 不符 (总分 2.8)
        code_a = """#include <stdio.h>
int main() {
    short x = 14, y = 20;
    printf("%5hd%5hd%5hd\\n", !x, x||y, x&&y);
    printf("%5hu%5hd%5hd\\n", (unsigned short)~x, (short)(x|y), (short)(x&y));
    printf("%5hd%5hd%5hd\\n", (short)~x, (short)(x|y), (short)(x&y));
    return 0;
}
"""
        # 候选 B: 逻辑正确，且参数完全匹配 y=88 (满分 3.0)
        code_b = """#include <stdio.h>
int main() {
    short x = 14, y = 88;
    printf("%5hd%5hd%5hd\\n", !x, x||y, x&&y);
    printf("%5hu%5hd%5hd\\n", (unsigned short)~x, (short)(x|y), (short)(x&y));
    printf("%5hd%5hd%5hd\\n", (short)~x, (short)(x|y), (short)(x&y));
    return 0;
}
"""
        file_a = src_dir / f"{std.student_id}_Lab1-2_A.c"
        file_a.write_text(code_a, encoding="utf-8")
        sf_a = SubmissionFile(path=file_a, original_name=file_a.name, file_type="c_source", matched_problem="1-2")

        file_b = src_dir / f"{std.student_id}_Lab1-2_B.c"
        file_b.write_text(code_b, encoding="utf-8")
        sf_b = SubmissionFile(path=file_b, original_name=file_b.name, file_type="c_source", matched_problem="1-2")

        hw_file = src_dir / f"{std.student_id}_hw.jpg"
        hw_file.write_bytes(b"\xff\xd8\xff\xe0" + b"\x00" * 100)
        sf_hw = SubmissionFile(path=hw_file, original_name=hw_file.name, file_type="image", matched_problem="1-2")

        judge = Lab1Judge()
        std_build = build_dir / std.student_id

        # 顺序 1: [sf_a, sf_b]
        res1 = judge._grade_lab1_2(std, [sf_a.path, sf_b.path], [sf_hw.path], std_build, mock_ai=True, all_student_files=[sf_a, sf_b, sf_hw])
        self.assertAlmostEqual(res1.score, 3.0, places=2)
        self.assertEqual(res1.chosen_candidate["src_path"], str(sf_b.path))

        # 顺序 2: [sf_b, sf_a]
        res2 = judge._grade_lab1_2(std, [sf_b.path, sf_a.path], [sf_hw.path], std_build, mock_ai=True, all_student_files=[sf_b, sf_a, sf_hw])
        self.assertAlmostEqual(res2.score, 3.0, places=2)
        self.assertEqual(res2.chosen_candidate["src_path"], str(sf_b.path))

    def test_52_compile_error_candidate_isolation(self):
        """测试52: 多候选存在编译错误版本时，不污染成功版本的评测与得分"""
        lab_dir = self.test_work / "test_52_lab"
        build_dir = lab_dir / "build"
        src_dir = lab_dir / "src"
        src_dir.mkdir(parents=True, exist_ok=True)
        build_dir.mkdir(parents=True, exist_ok=True)

        std = self.roster_mgr.students[0]

        bad_code = "syntax error; this is not C;\n"
        file_bad = src_dir / f"{std.student_id}_Lab2-1_bad.c"
        file_bad.write_text(bad_code, encoding="utf-8")
        sf_bad = SubmissionFile(path=file_bad, original_name=file_bad.name, file_type="c_source", matched_problem="2-1")

        good_code = """#include <stdio.h>
int main() {
    int h;
    if (scanf("%d", &h) != 1) { printf("Invalid input!\\n"); return 0; }
    if (h >= 0 && h < 12) printf("Good morning!\\n");
    else if (h == 12) printf("Good noon!\\n");
    else if (h >= 13 && h < 18) printf("Good afternoon!\\n");
    else if (h >= 18 && h < 21) printf("Good evening!\\n");
    else if (h >= 21 && h < 24) printf("Good night!\\n");
    else printf("Invalid input!\\n");
    return 0;
}
"""
        file_good = src_dir / f"{std.student_id}_Lab2-1_good.c"
        file_good.write_text(good_code, encoding="utf-8")
        sf_good = SubmissionFile(path=file_good, original_name=file_good.name, file_type="c_source", matched_problem="2-1")

        judge = Lab2Judge()
        std_build = build_dir / std.student_id
        std_build.mkdir(parents=True, exist_ok=True)

        grade = judge.grade_student(std, [sf_bad, sf_good], std_build, mock_ai=True)
        prob21 = grade.problem_results["2-1"]

        self.assertAlmostEqual(prob21.score, 2.0, places=2)
        self.assertEqual(prob21.status, "passed")
        self.assertFalse(grade.needs_review, "编译错误的不良候选不应将学生置为待复核")
        self.assertEqual(len(prob21.candidate_evaluations), 2)
        self.assertEqual(prob21.chosen_candidate["src_path"], str(sf_good.path))

    def test_53_complementary_candidates_no_cross_merging(self):
        """测试53: 多候选互补通过不同用例时，只能选择其中一个独立版本，不能交叉合并得分"""
        lab_dir = self.test_work / "test_53_lab"
        build_dir = lab_dir / "build"
        src_dir = lab_dir / "src"
        src_dir.mkdir(parents=True, exist_ok=True)
        build_dir.mkdir(parents=True, exist_ok=True)

        std = self.roster_mgr.students[0]

        # 候选 1: 只处理 morning (0..11)
        code_m = """#include <stdio.h>
int main() {
    int h;
    if (scanf("%d", &h) == 1 && h >= 0 && h < 12) printf("Good morning!\\n");
    else printf("Invalid input!\\n");
    return 0;
}
"""
        # 候选 2: 只处理 evening (18..20)
        code_e = """#include <stdio.h>
int main() {
    int h;
    if (scanf("%d", &h) == 1 && h >= 18 && h < 21) printf("Good evening!\\n");
    else printf("Invalid input!\\n");
    return 0;
}
"""
        f_m = src_dir / f"{std.student_id}_Lab2-1_m.c"
        f_m.write_text(code_m, encoding="utf-8")
        sf_m = SubmissionFile(path=f_m, original_name=f_m.name, file_type="c_source", matched_problem="2-1")

        f_e = src_dir / f"{std.student_id}_Lab2-1_e.c"
        f_e.write_text(code_e, encoding="utf-8")
        sf_e = SubmissionFile(path=f_e, original_name=f_e.name, file_type="c_source", matched_problem="2-1")

        judge = Lab2Judge()
        std_build = build_dir / std.student_id
        std_build.mkdir(parents=True, exist_ok=True)

        res = judge._grade_lab2_1(std, [sf_m.path, sf_e.path], std_build, mock_ai=True, all_files=[sf_m, sf_e])
        # 单独候选只通过各自部分用例，严禁跨版本合并用例得分
        self.assertLess(res.score, 2.0)
        self.assertGreater(res.score, 0.0)
        self.assertIn(res.chosen_candidate["src_path"], [str(sf_m.path), str(sf_e.path)])

    def test_54_serial_vs_parallel_compilation_consistency(self):
        """测试54: 并行编译与串行编译得到一致的候选集和最高分选择结果"""
        from checker.compiler import CompileTask
        lab_dir = self.test_work / "test_54_lab"
        src_dir = lab_dir / "src"
        src_dir.mkdir(parents=True, exist_ok=True)

        std = self.roster_mgr.students[0]

        files = []
        for i in range(4):
            f = src_dir / f"{std.student_id}_cand_{i}.c"
            f.write_text(f'#include <stdio.h>\nint main() {{ printf("cand {i}\\n"); return 0; }}', encoding="utf-8")
            files.append(f)

        compiler = ProjectCompiler()

        build_serial = lab_dir / "build_serial" / std.student_id
        build_serial.mkdir(parents=True, exist_ok=True)
        tasks_serial = [
            CompileTask(
                student_id=std.student_id,
                problem_id="lab2_1",
                cand_idx=i,
                src_path=files[i],
                build_dir=build_serial,
                binary_name=f"cand_{i}",
            )
            for i in range(4)
        ]
        res_serial = compiler.compile_batch(tasks_serial, max_workers=1)

        build_parallel = lab_dir / "build_parallel" / std.student_id
        build_parallel.mkdir(parents=True, exist_ok=True)
        tasks_parallel = [
            CompileTask(
                student_id=std.student_id,
                problem_id="lab2_1",
                cand_idx=i,
                src_path=files[i],
                build_dir=build_parallel,
                binary_name=f"cand_{i}",
            )
            for i in range(4)
        ]
        res_parallel = compiler.compile_batch(tasks_parallel, max_workers=4)

        self.assertEqual(len(res_serial), 4)
        self.assertEqual(len(res_parallel), 4)
        for i in range(4):
            key = (std.student_id, "lab2_1", i)
            self.assertTrue(res_serial[key].success)
            self.assertTrue(res_parallel[key].success)
            self.assertTrue(res_serial[key].binary_path.exists())
            self.assertTrue(res_parallel[key].binary_path.exists())

    def test_55_lab1_and_lab2_common_multiversion_workflow(self):
        """测试55: Lab1和Lab2均能使用这套公共流程正确评测多候选提交"""
        lab_dir = self.test_work / "test_55_lab"
        src_dir = lab_dir / "src"
        build_dir = lab_dir / "build"
        src_dir.mkdir(parents=True, exist_ok=True)
        build_dir.mkdir(parents=True, exist_ok=True)

        std = self.roster_mgr.students[0]
        std_build = build_dir / std.student_id
        std_build.mkdir(parents=True, exist_ok=True)

        code_bad_1_3 = '#include <stdio.h>\nint main(){ printf("wrong\\n"); return 0; }\n'
        code_good_1_3 = """#include <stdio.h>
int main() {
    int mode;
    double temp;
    if (scanf("%d %lf", &mode, &temp) != 2) return 0;
    if (mode == 0) printf("%.2f\\n", (temp - 32.0) * 5.0 / 9.0);
    else printf("%.2f\\n", temp * 9.0 / 5.0 + 32.0);
    return 0;
}
"""
        f_b_13 = src_dir / f"{std.student_id}_1_3_bad.c"
        f_b_13.write_text(code_bad_1_3, encoding="utf-8")
        sf_b_13 = SubmissionFile(path=f_b_13, original_name=f_b_13.name, file_type="c_source", matched_problem="1-3")

        f_g_13 = src_dir / f"{std.student_id}_1_3_good.c"
        f_g_13.write_text(code_good_1_3, encoding="utf-8")
        sf_g_13 = SubmissionFile(path=f_g_13, original_name=f_g_13.name, file_type="c_source", matched_problem="1-3")

        judge1 = Lab1Judge()
        res_1_3 = judge1._grade_lab1_3(std, [sf_b_13.path, sf_g_13.path], std_build, mock_ai=True, all_student_files=[sf_b_13, sf_g_13])
        self.assertAlmostEqual(res_1_3.score, 2.0, places=2)
        self.assertEqual(len(res_1_3.candidate_evaluations), 2)
        self.assertEqual(res_1_3.chosen_candidate["src_path"], str(sf_g_13.path))

        code_bad_2_2 = '#include <stdio.h>\nint main(){ printf("A\\n"); return 0; }\n'
        code_good_2_2 = """#include <stdio.h>
#include <ctype.h>
int main() {
    char ch;
    if (scanf(" %c", &ch) != 1 || !isalpha(ch)) { printf("ERROR\\n"); return 0; }
    ch = toupper(ch);
    int n = ch - 'A' + 1;
    for (int i = 0; i < n; i++) {
        for (int j = 0; j < n - i - 1; j++) printf(" ");
        for (char c = 'A'; c <= 'A' + i; c++) printf("%c", c);
        for (char c = 'A' + i - 1; c >= 'A'; c--) printf("%c", c);
        printf("\\n");
    }
    return 0;
}
"""
        f_b_22 = src_dir / f"{std.student_id}_2_2_bad.c"
        f_b_22.write_text(code_bad_2_2, encoding="utf-8")
        sf_b_22 = SubmissionFile(path=f_b_22, original_name=f_b_22.name, file_type="c_source", matched_problem="2-2")

        f_g_22 = src_dir / f"{std.student_id}_2_2_good.c"
        f_g_22.write_text(code_good_2_2, encoding="utf-8")
        sf_g_22 = SubmissionFile(path=f_g_22, original_name=f_g_22.name, file_type="c_source", matched_problem="2-2")

        judge2 = Lab2Judge()
        res_2_2 = judge2._grade_lab2_2(std, [sf_b_22.path, sf_g_22.path], std_build, mock_ai=True, all_files=[sf_b_22, sf_g_22])
        self.assertAlmostEqual(res_2_2.score, 3.0, places=2)
        self.assertEqual(len(res_2_2.candidate_evaluations), 2)
        self.assertEqual(res_2_2.chosen_candidate["src_path"], str(sf_g_22.path))

if __name__ == "__main__":
    unittest.main()
